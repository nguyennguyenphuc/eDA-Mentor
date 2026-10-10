# -*- coding: utf-8 -*-
"""Dựng _mentor/Nexa_Sales_P1_Doi_chieu_cong_thuc.xlsx — bản ĐỐI CHIẾU SỐ bằng công thức thường, 4 sheet giống HTML.
Dashboard chính (Power Query + Mô hình Dữ liệu + DAX) do build_dashboard_excel_com.py dựng trên Excel Windows.

- Tong_quan · Doanh_thu · Doi_soat · Hoan_tra: câu "Điểm chính", thẻ KPI (so tháng trước, cùng kỳ), tiêu đề hành động,
  biểu đồ (xám + một màu nhấn, đỏ chỉ cho vượt ngưỡng), bảng chi tiết — cùng bố cục và cùng số với HTML.
- Ô vàng ở sheet Tong_quan chọn Tháng / Kênh / Quốc gia; 4 sheet dùng chung lựa chọn đó.
- Data (tháng × cửa hàng), Data_dm (× danh mục), Data_ld (× lý do trả hàng): số tổng hợp từ aggregate.load()
  (đọc lại file giao học viên, cùng nguồn với HTML). Calc: mọi phép tính trung gian (ẩn).
- Chỉ dùng công thức thường (SUMIFS, COUNTIFS, INDEX/MATCH, LARGE, FIXED…) — không cần Power Pivot, mở được trên
  mọi Excel và LibreOffice. Số trong chuỗi dùng FIXED() để không phụ thuộc dấu thập phân của máy.
Đây là file ĐỐI CHIẾU SỐ cho mentor, không phải bài mẫu Power Pivot. Kiểm bằng: python xuat_pdf_kiem_tra.py
"""
import aggregate as A
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabel, DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = A.P1 / "_mentor" / "Nexa_Sales_P1_Doi_chieu_cong_thuc.xlsx"   # bản đối chiếu số; dashboard chính: build_dashboard_excel_com.py
REF = A.REF
D = A.load()
st = D["st"]
CAT_VI, COUNTRY_VI, REASON = D["CAT_VI"], D["COUNTRY_VI"], D["REASON"]
THR_ABS, THR_PCT = REF["threshold"]["abs_usd"], REF["threshold"]["pct"]
ASOF = f"{REF['as_of'][8:10]}/{REF['as_of'][5:7]}/{REF['as_of'][:4]}"

FONT = "Arial"
INK, INK2, MUTED, ACCENT, CTX, CTXD, DOWN, CRIT, RULE, UP = ("1F1F1E", "52514E", "8A8984", "2A78D6", "B9B7B0",
                                                             "8F8D86", "EB6834", "D03B3B", "E4E3DF", "1C5CAB")
f = lambda **k: Font(name=FONT, **{"size": 10, "color": INK, **k})
YELLOW = PatternFill("solid", fgColor="FFF2B3")
STORY = PatternFill("solid", fgColor="F2F6FC")
CARD = PatternFill("solid", fgColor="FAFAF8")
thin, bar = Side(style="thin", color=RULE), Side(style="thick", color=ACCENT)
NUM0, NUM1 = '#,##0;-#,##0;0', '#,##0.0;-#,##0.0;0.0'
PCT1, PCT2 = '0.0%', '0.00%'

wb = Workbook()
S1 = wb.active
S1.title = "Tong_quan"
S2, S3, S4 = wb.create_sheet("Doanh_thu"), wb.create_sheet("Doi_soat"), wb.create_sheet("Hoan_tra")
SH = wb.create_sheet("Huong_dan")
SD, SM, SL, SC = wb.create_sheet("Data"), wb.create_sheet("Data_dm"), wb.create_sheet("Data_ld"), wb.create_sheet("Calc")


def tbl(ws, name, head, rows):
    ws.append(head)
    for r in rows:
        ws.append(r)
    n = len(rows) + 1
    ws.add_table(Table(displayName=name, ref=f"A1:{CL(len(head))}{n}",
                       tableStyleInfo=TableStyleInfo(name="TableStyleLight15", showRowStripes=True)))
    ws.freeze_panes = "A2"
    for j, h in enumerate(head, 1):
        ws.column_dimensions[CL(j)].width = max(10, len(h) + 2)
    return n


# ------------------------------------------------------------------ dữ liệu (giá trị) + cột công thức theo lựa chọn
ms = D["ms"].merge(st[["store_key", "store_code", "store_name", "country", "channel_code"]], on="store_key")
ms = ms.sort_values(["month_id", "store_key"]).reset_index(drop=True)
HD = ["month_id", "month_label", "store_key", "store_code", "store_name", "country", "channel", "gross_sales",
      "discounts", "net_sales", "returns", "orders", "units", "finance_net_revenue", "net_sales_delivery_basis",
      "cancelled_orders", "pending_orders", "duplicate_lines", "delta", "over_threshold", "in_selection",
      "rank_net_sales", "rank_abs_delta", "rank_flag_pos"]
C = {h: CL(i + 1) for i, h in enumerate(HD)}
rows = []
for i, r in enumerate(ms.itertuples(), 2):
    rows.append([int(r.month_id), f"{r.month_id % 100:02d}/{r.month_id // 100}", int(r.store_key), r.store_code,
                 r.store_name, COUNTRY_VI[r.country], r.channel_code, round(r.gross, 2), round(r.disc, 2),
                 round(r.ns, 2), round(r.ret, 2), int(r.orders), int(r.units), round(r.fin, 2), round(r.ns_post, 2),
                 int(r.canc), int(r.pend), int(r.dup),
                 f"=N{i}-(J{i}-K{i})",
                 # AND() không dừng sớm: chia cho N=0 vẫn bị tính → tách IF như DIVIDE của DAX
                 f"=IF(N{i}=0,0,IF(AND(ABS(S{i})>Calc!$B$8,ABS(S{i})/ABS(N{i})>Calc!$B$9),1,0))",
                 f'=IF(AND(A{i}=Calc!$B$3,OR(Calc!$B$6="*",G{i}=Calc!$B$6),OR(Calc!$B$7="*",F{i}=Calc!$B$7)),1,0)',
                 f"=IF(AND(U{i}=1,J{i}>0),J{i}+ROW()/1000000,0)",
                 f"=IF(AND(U{i}=1,OR(N{i}<>0,J{i}<>0)),ABS(S{i})+ROW()/1000000,0)",
                 f"=IF(AND(U{i}=1,T{i}=1,C{i}<>0),ABS(S{i})+ROW()/1000000,0)"])
ND = tbl(SD, "tbl_Data", HD, rows)
for row in SD.iter_rows(min_row=2, min_col=8, max_col=19):
    for c in row:
        c.number_format = '#,##0.00'
R = lambda h: f"Data!${C[h]}$2:${C[h]}${ND}"

cat = D["cat"].merge(st[["store_key", "country", "channel_code"]], on="store_key")
NM = tbl(SM, "tbl_Data_dm", ["month_id", "store_key", "channel", "country", "category", "net_sales", "returns"],
         [[int(r.month_id), int(r.store_key), r.channel_code, COUNTRY_VI[r.country], CAT_VI[r.category],
           round(r.ns, 2), round(r.ret, 2)] for r in cat.itertuples()])
RM = lambda c: f"Data_dm!${c}$2:${c}${NM}"
rr = D["rr"].merge(st[["store_key", "country", "channel_code"]], on="store_key")
NL = tbl(SL, "tbl_Data_ld", ["month_id", "store_key", "channel", "country", "reason", "returns"],
         [[int(r.month_id), int(r.store_key), r.channel_code, COUNTRY_VI[r.country], REASON[r.return_reason],
           round(r.ret, 2)] for r in rr.itertuples()])
RL = lambda c: f"Data_ld!${c}$2:${c}${NL}"

# ------------------------------------------------------------------ Calc: tham số
months = sorted(m for m in ms.month_id.unique() if ms[ms.month_id == m].orders.sum() > 0)
labels = [f"{m % 100:02d}/{m // 100}" for m in months]
SC["A1"], SC["B1"] = "Tham số", "Giá trị"
P = [("Tháng chọn (nhãn)", "=Tong_quan!$C$5"),                                        # B2
     ("Tháng chọn (month_id)", f"=INDEX($U$2:$U${len(months) + 1},MATCH(B2,$T$2:$T${len(months) + 1},0))"),  # B3
     ("Tháng trước", "=IF(MOD(B3,100)=1,B3-89,B3-1)"),                                  # B4
     ("Cùng kỳ năm trước", "=B3-100"),                                                  # B5
     ("Kênh (tiêu chí)", '=IF(Tong_quan!$F$5="Cửa hàng (POS)","POS",IF(Tong_quan!$F$5="Website/app","ONLINE","*"))'),
     ("Quốc gia (tiêu chí)", '=IF(Tong_quan!$I$5="Tất cả","*",Tong_quan!$I$5)'),        # B7
     ("Ngưỡng USD", THR_ABS), ("Ngưỡng %", THR_PCT),                                    # B8, B9
     ("Phạm vi (chữ)", '=IF(B6="*","mọi kênh",IF(B6="POS","kênh cửa hàng","kênh online"))&IF(B7="*",""," · "&B7)'),
     ("Có kênh online trong phạm vi", '=IF(AND(B6<>"POS",B7="*"),1,0)')]                 # B11
for i, (k, v) in enumerate(P, 2):
    SC[f"A{i}"], SC[f"B{i}"] = k, v
# danh sách tra cứu đặt ở cột T:Z, tách khỏi các khối tính (cột A:R) để không tạo vòng tham chiếu
SC["T1"], SC["U1"] = "Danh sách tháng", "month_id"
for i, (lb, m) in enumerate(zip(reversed(labels), reversed(months)), 2):
    SC[f"T{i}"], SC[f"U{i}"] = lb, int(m)
SC["W1"] = "Quốc gia"
CTRY = sorted({COUNTRY_VI[c] for c in st[st.store_key > 0].country}, key=str)
for i, c in enumerate(["Tất cả"] + CTRY, 2):
    SC[f"W{i}"] = c
SC["Y1"], SC["Z1"] = "HQ-ADJ month_id", "USD"
for i, (m, v) in enumerate(D["hq"].items(), 2):
    SC[f"Y{i}"], SC[f"Z{i}"] = int(m), round(float(v), 2)

FILT = f'{R("channel")},$B$6,{R("country")},$B$7'


def su(col, mcell):                      # tổng một cột Data trong tháng mcell, theo kênh + quốc gia đang chọn
    return f"SUMIFS({R(col)},{R('month_id')},{mcell},{FILT})"


def cnt_flag(mcell, pos_only=False):
    extra = f',{R("store_key")},"<>0"' if pos_only else ""
    return f"COUNTIFS({R('month_id')},{mcell},{FILT},{R('over_threshold')},1{extra})"


# ------------------------------------------------------------------ Calc: KPI tháng chọn / LM / SPLY (hàng 15+)
K = {}
METRICS = [("gross", "gross_sales"), ("disc", "discounts"), ("ns", "net_sales"), ("ret", "returns"),
           ("orders", "orders"), ("units", "units"), ("fin", "finance_net_revenue"),
           ("nspost", "net_sales_delivery_basis"), ("canc", "cancelled_orders"), ("pend", "pending_orders")]
SC["A14"], SC["B14"], SC["C14"], SC["D14"] = "KPI", "Tháng chọn", "Tháng trước", "Cùng kỳ"
r0 = 15
for j, (k, col) in enumerate(METRICS):
    row = r0 + j
    SC[f"A{row}"] = k
    for cc, mc in zip("BCD", ("$B$3", "$B$4", "$B$5")):
        SC[f"{cc}{row}"] = "=" + su(col, mc)
    K[k] = row
row = r0 + len(METRICS)
DER = [("nr", "={c}{ns}-{c}{ret}"), ("aov", '=IF({c}{orders}=0,"",{c}{ns}/{c}{orders})'),
       ("rr", '=IF({c}{ns}=0,"",{c}{ret}/{c}{ns})'), ("delta", "={c}{fin}-{c}{nr}"),
       ("dadj", "={c}{fin}-({c}{nspost}-{c}{ret})"), ("disc_g", '=IF({c}{gross}=0,"",{c}{disc}/{c}{gross})')]
for k, fm in DER:
    SC[f"A{row}"] = k
    K[k] = row
    for cc in "BCD":
        SC[f"{cc}{row}"] = fm.format(c=cc, **K)
    row += 1
for k, fm in [("nflag", "=" + cnt_flag("$B$3")), ("nflagpos", "=" + cnt_flag("$B$3", True)),
              ("nrows", f'=COUNTIFS({R("month_id")},$B$3,{FILT},{R("rank_abs_delta")},">0")'),
              ("onl_ns", f'=IF($B$11=1,SUMIFS({R("net_sales")},{R("month_id")},$B$3,{R("store_key")},0),0)'),
              ("onl_nspost", f'=IF($B$11=1,SUMIFS({R("net_sales_delivery_basis")},{R("month_id")},$B$3,{R("store_key")},0),0)'),
              ("hq", "=SUMIFS($Z$2:$Z$10,$Y$2:$Y$10,$B$3)"),
              ("dup_all", f"=SUMIFS({R('duplicate_lines')},{FILT})"), ("pend_all", f"=SUMIFS({R('pending_orders')},{FILT})"),
              ("onl_ns_sp", f'=IF($B$11=1,SUMIFS({R("net_sales")},{R("month_id")},$B$5,{R("store_key")},0),0)'),
              ("pos_ns", f'=SUMIFS({R("net_sales")},{R("month_id")},$B$3,{R("channel")},"POS",{R("country")},$B$7)*($B$6<>"ONLINE")'),
              ("pos_ns_sp", f'=SUMIFS({R("net_sales")},{R("month_id")},$B$5,{R("channel")},"POS",{R("country")},$B$7)*($B$6<>"ONLINE")')]:
    SC[f"A{row}"], SC[f"B{row}"] = k, fm
    K[k] = row
    row += 1
SC[f"A{row}"], SC[f"B{row}"] = "timing", f"=B{K['onl_nspost']}-B{K['onl_ns']}"   # Δ gốc − Δ sau quy đổi (online)
K["timing"] = row
row += 1
k_ = lambda k, c="B": f"Calc!${c}${K[k]}"          # ô KPI để các sheet trỏ tới


def chg(k):                              # (% so tháng trước, % so cùng kỳ) dạng ô Calc
    out = []
    for c in "CD":
        r_ = row_chg[(k, c)]
        out.append(f"Calc!$F${r_}")
    return out


# % thay đổi (cột F) — số đo dạng tỷ lệ so bằng điểm %
row_chg = {}
SC["E14"], SC["F14"] = "Thay đổi", "Giá trị"
rr_ = r0
for k in ["nr", "orders", "aov", "rr", "ns", "units", "ret", "canc", "disc_g", "fin"]:
    for c in "CD":
        SC[f"E{rr_}"] = f"{k} vs {'LM' if c == 'C' else 'SPLY'}"
        if k in ("rr", "disc_g"):
            SC[f"F{rr_}"] = f'=IF(OR(B{K[k]}="",{c}{K[k]}=""),"",B{K[k]}-{c}{K[k]})'
        else:
            SC[f"F{rr_}"] = f'=IF(OR({c}{K[k]}=0,{c}{K[k]}=""),"",(B{K[k]}-{c}{K[k]})/ABS({c}{K[k]}))'
        row_chg[(k, c)] = rr_
        rr_ += 1


def sp(cell, d=1, pp=False):             # "+12,3%" / "−4,5%" / "–" (dùng FIXED: không phụ thuộc dấu thập phân của máy)
    if pp:
        return f'IF({cell}="","–",IF({cell}>=0,"+","−")&FIXED(ABS({cell})*100,{d})&" điểm %")'
    return f'IF({cell}="","–",IF({cell}>=0,"+","−")&FIXED(ABS({cell})*100,{d})&"%")'


def pc(cell, d=0):
    return f'IF({cell}="","–",FIXED({cell}*100,{d})&"%")'


def k1(cell):
    return f"FIXED({cell}/1000,1)"


def sg(cell):
    return f'IF({cell}>=0.5,"+",IF({cell}<=-0.5,"−",""))&FIXED(ABS({cell}),0)'


# ------------------------------------------------------------------ Calc: chuỗi 24 tháng (hàng 60–83)
W0, WN = 60, 24
WC = ["month_id", "label", "nr", "nr_sel", "orders", "orders_sel", "aov", "aov_sel", "dadj", "dadj_red", "dadj_gray",
      "rr", "rr_sel", "canc", "pend", "onl_raw", "onl_adj", "nflagpos"]
WX = {h: CL(i + 1) for i, h in enumerate(WC)}
SHOW = {"onl_raw": "Δ gốc", "onl_adj": "Sau quy đổi", "canc": "Huỷ", "pend": "Đang giao"}
for j, h in enumerate(WC):
    SC.cell(W0 - 1, j + 1, SHOW.get(h, h))
for i in range(WN):
    r = W0 + i
    m = f"$A{r}"
    SC[f"A{r}"] = f"=YEAR(DATE(INT($B$3/100),MOD($B$3,100)-{WN - 1 - i},1))*100+MONTH(DATE(INT($B$3/100),MOD($B$3,100)-{WN - 1 - i},1))"
    SC[f"B{r}"] = f'=RIGHT("0"&MOD(A{r},100),2)&"/"&RIGHT(INT(A{r}/100),2)'
    nr = f"({su('net_sales', m)}-{su('returns', m)})"
    has = f"COUNTIFS({R('month_id')},{m},{R('orders')},\">0\")>0"
    SC[f"C{r}"] = f"=IF({has},{nr}/1000,NA())"
    SC[f"E{r}"] = f"=IF({has},{su('orders', m)},NA())"
    SC[f"G{r}"] = f"=IF(AND({has},E{r}>0),{su('net_sales', m)}/E{r},NA())"
    SC[f"I{r}"] = f"=IF({has},({su('finance_net_revenue', m)}-({su('net_sales_delivery_basis', m)}-{su('returns', m)}))/1000,NA())"
    SC[f"R{r}"] = "=" + cnt_flag(m, True)
    SC[f"J{r}"] = f"=IF(ISNA(I{r}),0,IF(R{r}>0,I{r},0))"
    SC[f"K{r}"] = f"=IF(ISNA(I{r}),0,IF(R{r}>0,0,I{r}))"
    SC[f"L{r}"] = f"=IF({has},IF({su('net_sales', m)}=0,NA(),{su('returns', m)}/{su('net_sales', m)}*100),NA())"
    SC[f"N{r}"] = f"=IF({has},{su('cancelled_orders', m)},NA())"
    SC[f"O{r}"] = f"=IF({has},{su('pending_orders', m)},NA())"
    onl = lambda col: f'SUMIFS({R(col)},{R("month_id")},{m},{R("store_key")},0)'
    SC[f"P{r}"] = f"=IF($B$11=1,({onl('finance_net_revenue')}-({onl('net_sales')}-{onl('returns')}))/1000,NA())"
    SC[f"Q{r}"] = f"=IF($B$11=1,({onl('finance_net_revenue')}-({onl('net_sales_delivery_basis')}-{onl('returns')}))/1000,NA())"
    last = i == WN - 1
    for a, b in (("D", "C"), ("F", "E"), ("H", "G"), ("M", "L")):
        SC[f"{a}{r}"] = f"={b}{r}" if last else "=NA()"
WL = W0 + WN - 1
wref = lambda col: (W0, WL, WX[col])


# ------------------------------------------------------------------ Calc: xếp hạng (danh mục, quốc gia, lý do, cửa hàng)
def ranked(top, names, value_formula, label_formula=None, title=""):
    """Khối xếp hạng giảm dần: A tên, B giá trị, C khoá (giá trị + phân biệt), D tên sắp xếp, E giá trị sắp xếp,
    F nhấn (mục 1), G nền (còn lại), H nhãn. value_formula(r, name_cell) → công thức giá trị."""
    SC.cell(top - 1, 1, title)
    n = len(names)
    for i, nm in enumerate(names):
        r = top + i
        SC[f"A{r}"] = nm
        SC[f"B{r}"] = "=" + value_formula(r, f"$A{r}")
        SC[f"C{r}"] = f"=B{r}+ROW()/1000000000"
        SC[f"D{r}"] = f"=INDEX($A${top}:$A${top + n - 1},MATCH(LARGE($C${top}:$C${top + n - 1},{i + 1}),$C${top}:$C${top + n - 1},0))"
        SC[f"E{r}"] = f"=INDEX($B${top}:$B${top + n - 1},MATCH(LARGE($C${top}:$C${top + n - 1},{i + 1}),$C${top}:$C${top + n - 1},0))"
        SC[f"F{r}"] = f"=IF(ROW()={top},E{r},0)"
        SC[f"G{r}"] = f"=IF(ROW()={top},0,E{r})"
        if label_formula:
            SC[f"H{r}"] = "=" + label_formula(r)
    return top, top + n - 1


CATS = [CAT_VI[c] for c in sorted(CAT_VI)]
dm = lambda col, extra="": (f"SUMIFS({RM(col)},{RM('A')},$B$3,{RM('C')},$B$6,{RM('D')},$B$7{extra})")
NSB = f"$B${K['ns']}"
CAT_T = ranked(100, CATS, lambda r, nm: dm("F", f",{RM('E')},{nm}") + "/1000",
               lambda r: f'IF(E{r}=0,"",{k1(f"E{r}*1000")}&" · "&{pc(f"E{r}*1000/{NSB}")})', "Net Sales theo danh mục (nghìn USD)")
CTRY_ALL = CTRY + ["Online"]
CTY_T = ranked(115, CTRY_ALL, lambda r, nm: f"SUMIFS({R('net_sales')},{R('month_id')},$B$3,{FILT},{R('country')},{nm})/1000",
               lambda r: f'IF(E{r}=0,"",{k1(f"E{r}*1000")}&" · "&{pc(f"E{r}*1000/{NSB}")})', "Net Sales theo quốc gia (nghìn USD)")
RSN = [REASON[k] for k in sorted(REASON)]
RET_RSN = f"SUM($B$130:$B${130 + len(RSN) - 1})"
RET_RSN_X = f"SUM(Calc!$B$130:$B${130 + len(RSN) - 1})"
RSN_T = ranked(130, RSN, lambda r, nm: f"SUMIFS({RL('F')},{RL('A')},$B$3,{RL('C')},$B$6,{RL('D')},$B$7,{RL('E')},{nm})",
               lambda r: f'IF(E{r}=0,"",FIXED(E{r},0)&" · "&{pc(f"E{r}/{RET_RSN}")})', "Hàng bán bị trả lại theo lý do (USD)")
CRR_T = ranked(140, CATS, lambda r, nm: f'IF({dm("F", f",{RM("E")},{nm}")}=0,0,{dm("G", f",{RM("E")},{nm}")}/{dm("F", f",{RM("E")},{nm}")}*100)',
               lambda r: f'IF(E{r}=0,"",FIXED(E{r},1)&"% · NS "&FIXED(INDEX($B${CAT_T[0]}:$B${CAT_T[1]},MATCH(D{r},$A${CAT_T[0]}:$A${CAT_T[1]},0)),1)&"k")',
               "Tỷ lệ hàng bán bị trả lại theo danh mục (%)")

# Top 10 cửa hàng theo Net Sales (hàng 160+) và bảng đối soát top 10 theo |Δ| (hàng 175+)
TOP, REC = 160, 175
SC.cell(TOP - 1, 1, "Top 10 cửa hàng: dòng Data · tên · quốc gia · Net Sales · đơn · vs LM · vs SPLY")
for k in range(10):
    r = TOP + k
    SC[f"A{r}"] = f'=IF(LARGE({R("rank_net_sales")},{k + 1})>0,MATCH(LARGE({R("rank_net_sales")},{k + 1}),{R("rank_net_sales")},0),"")'
    ix = lambda col: f'IF($A{r}="","",INDEX({R(col)},$A{r}))'
    SC[f"B{r}"], SC[f"C{r}"], SC[f"D{r}"], SC[f"E{r}"] = "=" + ix("store_name"), "=" + ix("country"), "=" + ix("net_sales"), "=" + ix("orders")
    SC[f"F{r}"] = "=" + ix("store_key")
    for cc, mc in (("G", "$B$4"), ("H", "$B$5")):
        base = f'SUMIFS({R("net_sales")},{R("store_key")},F{r},{R("month_id")},{mc})'
        SC[f"{cc}{r}"] = f'=IF(OR($A{r}="",{base}=0),"",(D{r}-{base})/ABS({base}))'
SC.cell(REC - 1, 1, "Đối soát top 10 theo |Δ|: dòng · cửa hàng · đơn · Kế toán · Δ · Δ% · Δ sau quy đổi · cờ")
for k in range(10):
    r = REC + k
    SC[f"A{r}"] = f'=IF(LARGE({R("rank_abs_delta")},{k + 1})>0,MATCH(LARGE({R("rank_abs_delta")},{k + 1}),{R("rank_abs_delta")},0),"")'
    ix = lambda col: f'INDEX({R(col)},$A{r})'
    SC[f"B{r}"] = f'=IF($A{r}="","",{ix("store_code")}&" · "&{ix("country")})'
    SC[f"C{r}"] = f'=IF($A{r}="","",{ix("net_sales")}-{ix("returns")})'
    SC[f"D{r}"] = f'=IF($A{r}="","",{ix("finance_net_revenue")})'
    SC[f"E{r}"] = f'=IF($A{r}="","",{ix("delta")})'
    SC[f"F{r}"] = f'=IF(OR($A{r}="",D{r}=0),"",E{r}/ABS(D{r}))'
    SC[f"G{r}"] = f'=IF($A{r}="","",{ix("finance_net_revenue")}-({ix("net_sales_delivery_basis")}-{ix("returns")}))'
    SC[f"H{r}"] = f'=IF($A{r}="","",IF({ix("over_threshold")}=1,"⚠ Vượt","trong ngưỡng"))'

# Bảng kênh (hàng 190–191)
CHT = 190
for i, (code, name) in enumerate((("POS", "Cửa hàng"), ("ONLINE", "Online"))):
    r = CHT + i
    ok = f'AND(OR($B$6="*",$B$6="{code}"),OR("{code}"="POS",$B$7="*"))'
    s_ = lambda col, mc: f'SUMIFS({R(col)},{R("month_id")},{mc},{R("channel")},"{code}",{R("country")},$B$7)'
    SC[f"A{r}"], SC[f"B{r}"] = name, f"={ok}"
    SC[f"C{r}"] = f"=IF(B{r},{s_('net_sales', '$B$3')},\"\")"
    SC[f"D{r}"] = f"=IF(B{r},{s_('orders', '$B$3')},\"\")"
    SC[f"E{r}"] = f'=IF(AND(B{r},D{r}>0),C{r}/D{r},"")'
    for cc, mc in (("F", "$B$4"), ("G", "$B$5")):
        SC[f"{cc}{r}"] = f'=IF(OR(NOT(B{r}),{s_("net_sales", mc)}=0),"",(C{r}-{s_("net_sales", mc)})/ABS({s_("net_sales", mc)}))'


# Waterfall: 4 chuỗi chồng (nền dương ẩn, phần dương, nền âm ẩn, phần âm) cho từng bước a → b
def waterfall_block(top, steps):
    """steps: [(nhãn_công_thức, giá_trị_công_thức, là_tổng)] → cột A nhãn, B giá trị, C đầu, D cuối, E–H 4 chuỗi."""
    SC.cell(top - 1, 1, "Waterfall: nhãn · giá trị · từ · đến · nền+ · phần+ · nền− · phần−")
    for i, (lab, val, total) in enumerate(steps):
        r = top + i
        SC[f"A{r}"], SC[f"B{r}"] = lab, val
        SC[f"C{r}"] = "=0" if (total or i == 0) else f"=D{r - 1}"
        SC[f"D{r}"] = f"=B{r}" if total else f"=C{r}+B{r}"
        lo, hi = f"MIN(C{r},D{r})", f"MAX(C{r},D{r})"
        SC[f"E{r}"] = f"=MAX({lo},0)"
        SC[f"F{r}"] = f"=MAX({hi},0)-MAX({lo},0)"
        SC[f"G{r}"] = f"=MIN({hi},0)"
        SC[f"H{r}"] = f"=MIN({lo},0)-MIN({hi},0)"
        SC[f"I{r}"] = f'=IF(ABS(B{r})<0.5,"",IF({str(total).upper()},"",IF(B{r}>0,"+","−"))&FIXED(ABS(B{r}),1))'
    return top, top + len(steps) - 1


WF1 = waterfall_block(200, [('="Gross Sales"', f"={k_('gross')}/1000", True),
                            ('="Chiết khấu"', f"=-{k_('disc')}/1000", False),
                            ('="Hàng trả lại"', f"=-{k_('ret')}/1000", False),
                            ('="Net Revenue"', f"={k_('nr')}/1000", True)])
# phân rã Δ: online lệch thời điểm · 3 cửa hàng ⚠ lớn nhất · cửa hàng khác · Δ
FL = 210
SC.cell(FL - 1, 1, "3 cửa hàng POS vượt ngưỡng có |Δ| lớn nhất: dòng · mã · Δ")
for k in range(3):
    r = FL + k
    SC[f"A{r}"] = f'=IF(LARGE({R("rank_flag_pos")},{k + 1})>0,MATCH(LARGE({R("rank_flag_pos")},{k + 1}),{R("rank_flag_pos")},0),"")'
    SC[f"B{r}"] = f'=IF(A{r}="","—",INDEX({R("store_code")},A{r})&" ⚠")'
    SC[f"C{r}"] = f'=IF(A{r}="",0,INDEX({R("delta")},A{r}))'
WF2 = waterfall_block(220, [('="Online lệch thời điểm"', f"={k_('timing')}/1000", False)] +
                      [(f"=$B${FL + k}", f"=$C${FL + k}/1000", False) for k in range(3)] +
                      [('="Cửa hàng khác"', f"=({k_('delta')}-{k_('timing')}-SUM($C${FL}:$C${FL + 2}))/1000", False),
                       ('="Δ Kế toán − đơn"', f"={k_('delta')}/1000", True)])
SC.sheet_state = "hidden"
for c, w in zip("ABCDEFGHI", (24, 16, 14, 18, 14, 14, 14, 14, 12)):
    SC.column_dimensions[c].width = w


# ------------------------------------------------------------------ khung trang dashboard
CM = lambda ncol: round(ncol * 2.3 - 0.4, 1)     # bề rộng biểu đồ (cm) vừa ncol cột


def page(ws, title, n_cards):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.column_dimensions["A"].width = 2
    for j in range(2, 18):
        ws.column_dimensions[CL(j)].width = 11
    ws["B2"] = title
    ws["B2"].font = f(size=18, bold=True)
    ws["B3"] = (f"Nexa Home · Sales Project 1 · dữ liệu đến {ASOF} · nguồn: Nexa_Sales_P1_Dataset.xlsx, "
                "finance_revenue.xlsx (phòng Tài chính – Kế toán) · USD, chưa thuế GTGT")
    ws["B3"].font = f(size=9, color=MUTED)
    # dòng chọn: ô vàng ở Tong_quan, các sheet khác hiện lại lựa chọn
    for c, lab in (("B5", "Tháng báo cáo"), ("E5", "Kênh bán"), ("H5", "Quốc gia")):
        ws[c] = lab
        ws[c].font = f(color=INK2)
        ws[c].alignment = Alignment(horizontal="right")
    for c, src in (("C5", "C5"), ("F5", "F5"), ("I5", "I5")):
        if ws is not S1:
            ws[c] = f"=Tong_quan!{src}"
            ws[c].font = f(bold=True, size=11, color=INK2)
        ws[c].alignment = Alignment(horizontal="center")
    if ws is not S1:
        ws["K5"] = "← đổi lựa chọn ở sheet Tong_quan"
        ws["K5"].font = f(size=9, color=MUTED, italic=True)
    # dải Điểm chính
    ws["B7"] = f'="Điểm chính · "&Calc!$B$2&" · "&Calc!$B$10'
    ws["B7"].font = f(size=9, color=ACCENT, bold=True)
    for r, h in ((8, 34), (9, 18), (10, 18)):
        ws.merge_cells(f"B{r}:Q{r}")
        ws.row_dimensions[r].height = h
        ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="center")
        for j in range(2, 18):
            ws.cell(r, j).fill = STORY
        ws[f"B{r}"].border = Border(left=bar)
    ws["B8"].font = f(size=13, bold=True)
    ws["B9"].font = ws["B10"].font = f(size=10, color=INK2)
    for r in range(7, 11):
        ws[f"B{r}"].border = Border(left=bar)
    # thẻ KPI: mỗi thẻ 3 cột × 4 hàng (12–15)
    cards, wc = [], 15 // n_cards if n_cards == 5 else 16 // n_cards     # 5 thẻ × 3 cột hoặc 4 thẻ × 4 cột
    for i in range(n_cards):
        c0 = 2 + i * wc
        cards.append(c0)
        for r in range(12, 16):
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + wc - 1)
            for j in range(c0, c0 + wc):
                ws.cell(r, j).fill = CARD
                ws.cell(r, j).border = Border(top=thin if r == 12 else None, bottom=thin if r == 15 else None,
                                              left=thin if j == c0 else None, right=thin if j == c0 + wc - 1 else None)
        ws.cell(12, c0).font = f(size=9, color=INK2)
        ws.cell(13, c0).font = f(size=18, bold=True)
        ws.cell(14, c0).font = ws.cell(15, c0).font = f(size=9, color=INK2)
    ws.row_dimensions[13].height = 28
    ws["B60"] = ("Δ = số Kế toán − dữ liệu đơn; vượt ngưỡng khi |Δ| > " + f"{THR_ABS:,.0f} USD và |Δ %| > "
                 f"{THR_PCT * 100:.0f}% ở mức tháng × cửa hàng. Chỉ tính đơn COMPLETED, đã loại dòng trùng; hàng trả lại "
                 "tính theo ngày trả. Kênh online không thuộc quốc gia nào. File đối chiếu số – công thức thường, không dùng "
                 "Mô hình Dữ liệu.")
    ws["B60"].font = f(size=8, color=MUTED)
    ws.print_area = "A1:Q61"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    return cards


def card(ws, c0, label, value, fmt, comps=(), notes=()):
    """comps: [(ô thay đổi, chữ, tăng_là_tốt, điểm_%)]; notes: [công thức/chữ]."""
    ws.cell(12, c0, label)
    v = ws.cell(13, c0, value)
    v.number_format = fmt
    v.alignment = Alignment(horizontal="left")
    def so_sanh(cell, lab, pp):
        so = f'FIXED(ABS({cell})*100,2)&" điểm %"' if pp else f'FIXED(ABS({cell})*100,1)&"%"'
        return f'=IF({cell}="","– {lab}",IF({cell}>0,"▲ ",IF({cell}<0,"▼ ","■ "))&{so}&" {lab}")'
    lines = [so_sanh(cell, lab, pp) for cell, lab, good, pp in comps] + list(notes)
    for i, t in enumerate(lines[:2]):
        ws.cell(14 + i, c0, t)
    for i, (cell, lab, good, pp) in enumerate(comps[:2]):
        ref = f"{CL(c0)}{14 + i}"
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'LEFT({ref},1)="▲"'], font=Font(color=UP if good else DOWN, bold=True)))
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'LEFT({ref},1)="▼"'], font=Font(color=DOWN if good else UP, bold=True)))


def title(ws, cell, formula, sub):
    ws[cell] = formula
    ws[cell].font = f(size=11, bold=True)
    r = int(cell[1:])
    ws[f"{cell[0]}{r + 1}"] = sub
    ws[f"{cell[0]}{r + 1}"].font = f(size=8, color=MUTED)


def _axis_font(ax, size=800, color=MUTED):
    ax.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=size, solidFill=color)),
                                    endParaRPr=CharacterProperties(sz=size))])


def style(ch, w, h, legend=False):
    ch.width, ch.height = w, h
    if not legend:
        ch.legend = None
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="ECEBE7"))
    for ax in (ch.x_axis, ch.y_axis):
        _axis_font(ax)
        ax.delete = False
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=RULE))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    return ch


def series_color(s, color, line=False, width=22000, nofill=False):
    if nofill:
        s.graphicalProperties.noFill = True
        s.graphicalProperties.line.noFill = True
    elif line:
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = width
        s.smooth = False
    else:
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.solidFill = color


def _txt(color=INK, size=800):
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=size, b=True, solidFill=color)),
                                 endParaRPr=CharacterProperties(sz=size))])


def labels(s, fmt="#,##0;-#,##0;;", pos=None, color=INK):
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.numFmt = fmt
    for k in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
        setattr(s.dLbls, k, False)
    if pos:
        s.dLbls.position = pos
    s.dLbls.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=800, b=True, solidFill=color)),
                                         endParaRPr=CharacterProperties(sz=800))])


def line_chart(cols, w=17.2, h=7.2, sel=True, colors=(CTX,), fmt="#,##0"):
    ch = LineChart()
    cats = Reference(SC, min_col=2, min_row=W0, max_row=WL)
    for k, (col, color) in enumerate(zip(cols, colors)):
        c = ord(WX[col]) - 64
        ch.add_data(Reference(SC, min_col=c, min_row=W0 - 1, max_row=WL), titles_from_data=True)
        s = ch.series[-1]
        series_color(s, color, line=True)
        s.marker = Marker(symbol="none")
        if sel and k == len(cols) - 1:          # tháng đang chọn = điểm cuối cửa sổ 24 tháng
            pt = DataPoint(idx=WN - 1)
            pt.marker = Marker(symbol="circle", size=8)
            pt.marker.graphicalProperties = GraphicalProperties(solidFill=ACCENT, ln=LineProperties(solidFill=ACCENT))
            s.dPt.append(pt)
            s.dLbls = DataLabelList(dLbl=[DataLabel(idx=WN - 1, showVal=True, showSerName=False, showCatName=False,
                                                    showLegendKey=False, showPercent=False, numFmt=fmt, dLblPos="t",
                                                    txPr=_txt(ACCENT))])
            for k_ in ("showVal", "showSerName", "showCatName", "showLegendKey", "showPercent"):
                setattr(s.dLbls, k_, False)
    ch.set_categories(cats)
    ch.y_axis.number_format = fmt
    ch.x_axis.tickLblSkip = 3
    return style(ch, w, h, legend=len(cols) > 1)


def stacked_bar(block, cols, colors, horizontal=True, w=11.4, h=8.6, lab_col=None, fmt="#,##0.0"):
    top, bot = block
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.grouping, ch.overlap, ch.gapWidth = "stacked", 100, 60
    for col, color in zip(cols, colors):
        c = ord(col) - 64
        ch.add_data(Reference(SC, min_col=c, min_row=top, max_row=bot), titles_from_data=False)
        s = ch.series[-1]
        if color is None:
            series_color(s, None, nofill=True)
        else:
            series_color(s, color)
            labels(s, f"{fmt};-{fmt};;")
    ch.set_categories(Reference(SC, min_col=ord(lab_col or "D") - 64, min_row=top, max_row=bot))
    ch.y_axis.number_format = fmt
    style(ch, w, h)
    if not horizontal:
        ch.x_axis.tickLblSkip = 3
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"
        ch.y_axis.delete = True
        ch.y_axis.majorGridlines = None
    return ch


def waterfall_chart(block, colors, w=11.4, h=7.2):
    top, bot = block
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap, ch.gapWidth = "col", "stacked", 100, 50
    for col, vis in (("E", False), ("F", True), ("G", False), ("H", True)):
        ch.add_data(Reference(SC, min_col=ord(col) - 64, min_row=top, max_row=bot), titles_from_data=False)
        s = ch.series[-1]
        if not vis:
            series_color(s, None, nofill=True)
            continue
        series_color(s, CTXD)
        for i, c in enumerate(colors):
            pt = DataPoint(idx=i)
            pt.graphicalProperties.solidFill = c
            pt.graphicalProperties.line.solidFill = c
            s.dPt.append(pt)
        labels(s, "#,##0.0;-#,##0.0;;")
    ch.set_categories(Reference(SC, min_col=1, min_row=top, max_row=bot))
    ch.y_axis.number_format = "#,##0"
    return style(ch, w, h)


def table_block(ws, r0, c0, head, src_rows, fmts, widths=None):
    for j, h in enumerate(head):
        c = ws.cell(r0, c0 + j, h)
        c.font = f(size=9, bold=True, color=INK2)
        c.border = Border(bottom=thin)
        c.alignment = Alignment(horizontal="left" if j == 0 or fmts[j] == "@" else "right")
    for i, cells in enumerate(src_rows):
        for j, v in enumerate(cells):
            c = ws.cell(r0 + 1 + i, c0 + j, v)
            c.font = f(size=9)
            c.number_format = fmts[j]
            c.alignment = Alignment(horizontal="left" if fmts[j] == "@" else "right")
            c.border = Border(bottom=Side(style="hair", color=RULE))


DPCT = '+0.0%;-0.0%;0.0%'
# ================================================================== Trang 1 · Tổng quan
cards = page(S1, "1 · Tổng quan: tháng này bán được bao nhiêu, có khớp Kế toán không?", 5)
S1["C5"], S1["F5"], S1["I5"] = f"{REF['report_month'] % 100:02d}/{REF['report_month'] // 100}", "Tất cả", "Tất cả"
for c in ("C5", "F5", "I5"):
    S1[c].fill, S1[c].font = YELLOW, f(bold=True, size=11)
S1.column_dimensions["C"].width = S1.column_dimensions["F"].width = S1.column_dimensions["I"].width = 14
S1["K5"] = "← ô vàng: chọn từ danh sách; cả 4 sheet tự cập nhật"
S1["K5"].font = f(size=9, color=MUTED, italic=True)
for cell, src in (("C5", f"Calc!$T$2:$T${len(months) + 1}"), ("F5", '"Tất cả,Cửa hàng (POS),Website/app"'),
                  ("I5", f"Calc!$W$2:$W${len(CTRY) + 2}")):
    dv = DataValidation(type="list", formula1=src, allow_blank=False)
    S1.add_data_validation(dv)
    dv.add(cell)
cL = lambda k: f"Calc!$F${row_chg[(k, 'C')]}"
cS = lambda k: f"Calc!$F${row_chg[(k, 'D')]}"
S1["B8"] = (f'="Net Revenue "&Calc!$B$2&" đạt "&{k1(k_("nr"))}&" nghìn USD — "&{sp(cL("nr"))}&" so với tháng trước, "'
            f'&{sp(cS("nr"))}&" so với cùng kỳ năm trước."')
S1["B9"] = (f'=IF(OR({cS("orders")}="",{cS("aov")}=""),"",IF(ABS(LN(1+{cS("orders")}))>=ABS(LN(1+{cS("aov")})),'
            f'"• Biến động so với cùng kỳ đến chủ yếu từ số đơn ("&{sp(cS("orders"))}&"); giá trị mỗi đơn (AOV) "&{sp(cS("aov"))}&".",'
            f'"• Biến động so với cùng kỳ đến chủ yếu từ giá trị mỗi đơn (AOV "&{sp(cS("aov"))}&"); số đơn "&{sp(cS("orders"))}&"."))')
S1["B10"] = (f'="• Kế toán ghi nhận "&IF({k_("delta")}>=0,"cao","thấp")&" hơn dữ liệu đơn "&FIXED(ABS({k_("delta")}),0)'
             f'&" USD ("&{sp(f"IF({k_("nr")}=0,\"\",{k_("delta")}/ABS({k_("nr")}))")}&"); "&{k_("nflag")}&" cửa hàng vượt ngưỡng cần giải thích (sheet Doi_soat)."')
card(S1, cards[0], "Net Revenue", f"={k_('nr')}/1000", '#,##0.0" nghìn USD"',
     [(cL("nr"), "so với tháng trước", True, False), (cS("nr"), "so với cùng kỳ", True, False)])
card(S1, cards[1], "Số đơn hoàn tất", f"={k_('orders')}", '#,##0" đơn"',
     [(cL("orders"), "so với tháng trước", True, False), (cS("orders"), "so với cùng kỳ", True, False)])
card(S1, cards[2], "AOV", f"={k_('aov')}", '#,##0" USD/đơn"',
     [(cL("aov"), "so với tháng trước", True, False), (cS("aov"), "so với cùng kỳ", True, False)])
card(S1, cards[3], "Tỷ lệ hàng bán bị trả lại", f"={k_('rr')}", PCT2,
     [(cL("rr"), "so với tháng trước", False, True), (cS("rr"), "so với cùng kỳ", False, True)])
card(S1, cards[4], "Chênh lệch với Kế toán (Δ)", f"={k_('delta')}", '+#,##0" USD";-#,##0" USD";0" USD"', (),
     [f'={sp(f"IF({k_("nr")}=0,\"\",{k_("delta")}/ABS({k_("nr")}))", 2)}&" so với dữ liệu đơn"',
      f'={k_("nflag")}&" cửa hàng vượt ngưỡng"'])
S1.conditional_formatting.add(f"{CL(cards[4])}13", FormulaRule(formula=[f"{k_('nflag')}>0"], font=Font(color=CRIT, bold=True)))
title(S1, "B17", f'="Net Revenue "&Calc!$B$2&" bằng "&{pc(f"IF({k_("nr", "D")}=0,\"\",{k_("nr")}/{k_("nr", "D")})")}&" cùng kỳ năm trước"',
      "Net Revenue theo tháng đặt hàng, nghìn USD · tháng đang chọn tô xanh")
S1.add_chart(line_chart(["nr"], w=CM(10), h=8, fmt="#,##0"), "B19")
title(S1, "L17", f'="Chiết khấu và hàng trả lại lấy đi "&{pc(f"IF({k_("gross")}=0,\"\",({k_("disc")}+{k_("ret")})/{k_("gross")})", 1)}&" Gross Sales"',
      "Tháng đang chọn, nghìn USD")
S1.add_chart(waterfall_chart(WF1, [CTXD, DOWN, DOWN, ACCENT], w=CM(6), h=8), "L19")
title(S1, "B35", f'="Số đơn "&{sp(cS("orders"))}&" so với cùng kỳ"', "Số đơn hoàn tất theo tháng")
S1.add_chart(line_chart(["orders"], w=CM(6), h=7, colors=(ACCENT,), fmt="#,##0"), "B37")
title(S1, "H35", f'="AOV "&{sp(cS("aov"))}&" so với cùng kỳ"', "AOV = Net Sales / số đơn, USD")
S1.add_chart(line_chart(["aov"], w=CM(6), h=7, fmt="#,##0"), "H37")
title(S1, "N35", f'="Sau quy đổi online, Δ "&Calc!$B$2&" còn "&{sg(k_("dadj"))}&" USD"',
      "Δ sau quy đổi theo tháng, nghìn USD · đỏ = có cửa hàng vượt ngưỡng")
S1.add_chart(stacked_bar((W0, WL), ["J", "K"], [CRIT, CTX], horizontal=False, w=CM(4), h=7, lab_col="B", fmt="#,##0.0"), "N37")
# bỏ nhãn số trên cột tháng (24 cột, chật): chỉ giữ màu
_ch = S1._charts[-1]
for s in _ch.series:
    s.dLbls = None

# ================================================================== Trang 2 · Doanh thu
cards = page(S2, "2 · Doanh thu đến từ đâu: kênh, danh mục, quốc gia, cửa hàng", 4)
top_cat, top_cty = f"Calc!$D${CAT_T[0]}", f"Calc!$D${CTY_T[0]}"
S2["B8"] = (f'=IF({k_("ns")}=0,"Không có doanh thu trong phạm vi đang chọn.",{top_cat}&" dẫn đầu với "&'
            f'{pc(f"Calc!$E${CAT_T[0]}*1000/{k_("ns")}")}&" Net Sales "&Calc!$B$2&"; "&{top_cty}&" là thị trường lớn nhất ("&'
            f'{pc(f"Calc!$E${CTY_T[0]}*1000/{k_("ns")}")}&").")')
S2["B9"] = (f'=IF(Calc!$B$6="*","• Online chiếm "&{pc(f"IF({k_("ns")}=0,\"\",{k_("onl_ns")}/{k_("ns")})")}&" Net Sales; so với cùng kỳ online "&'
            f'{sp(f"IF({k_("onl_ns_sp")}=0,\"\",({k_("onl_ns")}-{k_("onl_ns_sp")})/{k_("onl_ns_sp")})")}&", cửa hàng "&'
            f'{sp(f"IF({k_("pos_ns_sp")}=0,\"\",({k_("pos_ns")}-{k_("pos_ns_sp")})/{k_("pos_ns_sp")})")}&".","")')
S2["B10"] = f'="• Chiết khấu bằng "&{pc(k_("disc_g"), 1)}&" Gross Sales tháng này."'
card(S2, cards[0], "Net Sales", f"={k_('ns')}/1000", '#,##0.0" nghìn USD"',
     [(cL("ns"), "so với tháng trước", True, False), (cS("ns"), "so với cùng kỳ", True, False)])
card(S2, cards[1], "Số sản phẩm bán", f"={k_('units')}", '#,##0" sản phẩm"',
     [(cL("units"), "so với tháng trước", True, False), (cS("units"), "so với cùng kỳ", True, False)])
card(S2, cards[2], "Tỷ trọng online", f'=IF({k_("ns")}=0,"",{k_("onl_ns")}/{k_("ns")})', "0%", (), ["phần Net Sales từ website/app"])
card(S2, cards[3], "Chiết khấu / Gross Sales", f"={k_('disc_g')}", PCT1, [(cS("disc_g"), "so với cùng kỳ", False, True)])
title(S2, "B17", f'={top_cat}&" chiếm "&{pc(f"Calc!$E${CAT_T[0]}*1000/{k_("ns")}")}&" Net Sales"',
      "Net Sales tháng đang chọn theo danh mục, nghìn USD")
S2.add_chart(stacked_bar(CAT_T, ["F", "G"], [ACCENT, CTX], w=CM(6), h=8.6), "B19")
title(S2, "H17", f'={top_cty}&" đóng góp "&{pc(f"Calc!$E${CTY_T[0]}*1000/{k_("ns")}")}&" Net Sales"',
      "Net Sales theo quốc gia, nghìn USD (online tính riêng)")
S2.add_chart(stacked_bar(CTY_T, ["F", "G"], [ACCENT, CTX], w=CM(6), h=8.6), "H19")
title(S2, "N17", f'="AOV cửa hàng "&IF(Calc!$E${CHT}="","–",FIXED(Calc!$E${CHT},0))&" USD, online "&IF(Calc!$E${CHT + 1}="","–",FIXED(Calc!$E${CHT + 1},0))&" USD"',
      "So với tháng trước (LM) và cùng kỳ (SPLY)")
table_block(S2, 19, 14, ["Kênh", "Net Sales", "Đơn", "AOV"], [[f"=Calc!A{CHT + i}", f"=Calc!C{CHT + i}", f"=Calc!D{CHT + i}",
                                                               f"=Calc!E{CHT + i}"] for i in range(2)], ["@", NUM0, NUM0, NUM0])
table_block(S2, 23, 14, ["Kênh", "vs LM", "vs SPLY", ""], [[f"=Calc!A{CHT + i}", f"=Calc!F{CHT + i}", f"=Calc!G{CHT + i}", ""]
                                                          for i in range(2)], ["@", DPCT, DPCT, "@"])
title(S2, "B37", f'="10 cửa hàng lớn nhất tạo "&{pc(f"IF({k_("ns")}=0,\"\",SUM(Calc!$D${TOP}:$D${TOP + 9})/{k_("ns")})")}&" Net Sales"',
      "Top 10 cửa hàng theo Net Sales tháng đang chọn")
table_block(S2, 39, 2, ["#", "Cửa hàng", "", "", "Quốc gia", "", "Net Sales", "", "Đơn", "vs LM", "vs SPLY"],
            [[k + 1, f"=Calc!B{TOP + k}", "", "", f"=Calc!C{TOP + k}", "", f"=Calc!D{TOP + k}", "", f"=Calc!E{TOP + k}",
              f"=Calc!G{TOP + k}", f"=Calc!H{TOP + k}"] for k in range(10)],
            ["0", "@", "@", "@", "@", "@", NUM0, "@", NUM0, DPCT, DPCT])
S2.conditional_formatting.add("H40:H49", DataBarRule(start_type="num", start_value=0, end_type="max", color=ACCENT))

# ================================================================== Trang 3 · Đối soát
cards = page(S3, "3 · Đối soát với Kế toán: lệch bao nhiêu, ở đâu, vì sao?", 5)
S3["B8"] = (f'="Kế toán "&IF({k_("delta")}>=0,"cao","thấp")&" hơn dữ liệu đơn "&FIXED(ABS({k_("delta")}),0)&" USD ("&'
            f'{sp(f"IF({k_("fin")}=0,\"\",{k_("delta")}/ABS({k_("fin")}))", 2)}&")."&IF(ABS({k_("timing")})>1," "&{sg(k_("timing"))}&'
            f'" USD là lệch thời điểm của kênh online — biến mất khi tính theo ngày giao.","")')
S3["B9"] = (f'=IF({k_("nflagpos")}>0,"• Còn lại "&{sg(k_("dadj"))}&" USD; cần điều tra "&{k_("nflagpos")}&" cửa hàng vượt ngưỡng: "&'
            f'Calc!$B${FL}&IF(Calc!$A${FL + 1}<>"",", "&Calc!$B${FL + 1},"")&IF(Calc!$A${FL + 2}<>"",", "&Calc!$B${FL + 2},"")&".",'
            f'"• Còn lại "&{sg(k_("dadj"))}&" USD, không cửa hàng nào vượt ngưỡng ngoài kênh online.")')
S3["B10"] = (f'=IF({k_("hq")}<>0,"• Chưa gồm bút toán điều chỉnh HQ-ADJ "&{sg(k_("hq"))}&" USD (không thuộc cửa hàng nào).",'
             f'"• Xuống mức ngày ghi sổ (posting_date) ở các cửa hàng có ⚠ để tìm nguyên nhân và số tiền.")')
card(S3, cards[0], "Net Revenue · dữ liệu đơn", f"={k_('nr')}/1000", '#,##0.0" nghìn USD"')
card(S3, cards[1], "Net Revenue · Kế toán", f"={k_('fin')}/1000", '#,##0.0" nghìn USD"', (),
     [f'=IF({k_("hq")}<>0,"chưa gồm HQ-ADJ "&{sg(k_("hq"))}&" USD","")'])
card(S3, cards[2], "Chênh lệch Δ", f"={k_('delta')}", '+#,##0" USD";-#,##0" USD";0" USD"', (),
     [f'={sp(f"IF({k_("fin")}=0,\"\",{k_("delta")}/ABS({k_("fin")}))", 2)}&" trên số Kế toán"'])
card(S3, cards[3], "Δ sau quy đổi online", f"={k_('dadj')}", '+#,##0" USD";-#,##0" USD";0" USD"', (), ["online tính theo ngày giao"])
card(S3, cards[4], "Cửa hàng vượt ngưỡng", f'={k_("nflag")}&" / "&{k_("nrows")}', "@", (),
     [f"|Δ| > {THR_ABS:,.0f} USD và |Δ%| > {THR_PCT * 100:.0f}%"])
for c0, cond in ((cards[2], f"ABS({k_('delta')})>{THR_ABS}"), (cards[3], f"{k_('nflagpos')}>0"), (cards[4], f"{k_('nflag')}>0")):
    S3.conditional_formatting.add(f"{CL(c0)}13", FormulaRule(formula=[cond], font=Font(color=CRIT, bold=True)))
title(S3, "B17", f'=IF(Calc!$A${FL}<>"",INDEX({R("store_code")},Calc!$A${FL})&" là chênh lệch thật lớn nhất ("&{sg(f"Calc!$C${FL}")}&" USD)","Chênh lệch chủ yếu là lệch thời điểm")',
      "Phân rã Δ tháng đang chọn, nghìn USD · không gồm HQ-ADJ")
S3.add_chart(waterfall_chart(WF2, [CTX, CRIT, CRIT, CRIT, CTXD, ACCENT], w=CM(6), h=8), "B19")
title(S3, "H17", f'="Δ sau quy đổi "&Calc!$B$2&" là "&{sg(k_("dadj"))}&" USD"',
      "Δ sau khi quy online về ngày giao, nghìn USD · tháng có cửa hàng vượt ngưỡng tô đỏ")
S3.add_chart(stacked_bar((W0, WL), ["J", "K"], [CRIT, CTX], horizontal=False, w=CM(10), h=8, lab_col="B"), "H19")
for s in S3._charts[-1].series:
    s.dLbls = None
title(S3, "B35", f'=IF({k_("nflag")}>0,{k_("nflag")}&" cửa hàng vượt ngưỡng tháng "&Calc!$B$2,"Không cửa hàng nào vượt ngưỡng tháng "&Calc!$B$2)',
      "Xếp theo |Δ| giảm dần · ⚠ = |Δ| > ngưỡng USD và |Δ %| > ngưỡng %")
table_block(S3, 37, 2, ["Cửa hàng", "", "Đơn hàng", "Kế toán", "Δ", "Δ %", "Δ sau quy đổi", "Cờ"],
            [[f"=Calc!B{REC + k}", "", f"=Calc!C{REC + k}", f"=Calc!D{REC + k}", f"=Calc!E{REC + k}", f"=Calc!F{REC + k}",
              f"=Calc!G{REC + k}", f"=Calc!H{REC + k}"] for k in range(10)],
            ["@", "@", NUM0, NUM0, '+#,##0;-#,##0;0', DPCT, '+#,##0;-#,##0;0', "@"])
S3.conditional_formatting.add("B38:I47", FormulaRule(formula=['$I38="⚠ Vượt"'], font=Font(color=CRIT, bold=True)))
title(S3, "K35", '=IF(Calc!$B$11=1,"Δ online về gần 0 khi tính theo ngày giao","Kênh online không thuộc bộ lọc đang chọn")',
      "Δ kênh online theo tháng, nghìn USD · xám = gốc, xanh = sau quy đổi")
S3.add_chart(line_chart(["onl_raw", "onl_adj"], w=CM(7), h=7, colors=(CTX, ACCENT), fmt="#,##0.0"), "K37")

# ================================================================== Trang 4 · Hoàn trả & chất lượng
cards = page(S4, "4 · Hàng bán bị trả lại và chất lượng dữ liệu", 4)
S4["B8"] = (f'="Hàng trả lại "&Calc!$B$2&" là "&FIXED({k_("ret")},0)&" USD, bằng "&{pc(k_("rr"), 2)}&" Net Sales ("&'
            f'{sp(cS("rr"), 2, True)}&" so với cùng kỳ)."')
S4["B9"] = (f'=IF(Calc!$E${RSN_T[0]}=0,"• Không có hàng trả lại trong phạm vi đang chọn.","• Lý do lớn nhất: "&Calc!$D${RSN_T[0]}&" ("&'
            f'{pc(f"Calc!$E${RSN_T[0]}/{RET_RSN_X}")}&" giá trị trả).")')
S4["B10"] = (f'="• Trước khi tính KPI phải loại "&FIXED({k_("dup_all")},0)&" dòng trùng và "&FIXED({k_("pend_all")},0)&'
             f'" đơn đang giao (PENDING) — toàn thời gian, theo bộ lọc kênh/quốc gia."')
card(S4, cards[0], "Giá trị hàng trả lại", f"={k_('ret')}", '#,##0" USD"',
     [(cL("ret"), "so với tháng trước", False, False), (cS("ret"), "so với cùng kỳ", False, False)])
card(S4, cards[1], "Tỷ lệ hàng bán bị trả lại", f"={k_('rr')}", PCT2,
     [(cL("rr"), "so với tháng trước", False, True), (cS("rr"), "so với cùng kỳ", False, True)])
card(S4, cards[2], "Đơn huỷ", f"={k_('canc')}", '#,##0" đơn"', [(cS("canc"), "so với cùng kỳ", False, False)])
card(S4, cards[3], "Dòng trùng đã loại", f"={k_('dup_all')}", '#,##0" dòng"', (),
     ["toàn thời gian, theo bộ lọc", f'=FIXED({k_("pend_all")},0)&" đơn PENDING tại {ASOF}"'])
title(S4, "B17", f'="Tỷ lệ hàng trả lại "&{pc(k_("rr"), 2)}&" — "&IF({k_("rr")}>{k_("rr", "D")},"cao","thấp")&" hơn cùng kỳ "&'
                 f'IF({cS("rr")}="","–",FIXED(ABS({cS("rr")})*100,2))&" điểm %"', "Returns / Net Sales theo tháng, %")
S4.add_chart(line_chart(["rr"], w=CM(10), h=8, fmt="0.0"), "B19")
title(S4, "L17", f'=IF(Calc!$E${RSN_T[0]}=0,"Lý do trả hàng",Calc!$D${RSN_T[0]}&" là lý do trả hàng lớn nhất")',
      "Giá trị trả tháng đang chọn theo lý do, USD")
S4.add_chart(stacked_bar(RSN_T, ["F", "G"], [ACCENT, CTX], w=CM(6), h=8, fmt="#,##0"), "L19")
title(S4, "B35", f'=Calc!$D${CRR_T[0]}&" có tỷ lệ trả hàng cao nhất ("&FIXED(Calc!$E${CRR_T[0]},1)&"%)"',
      "Tỷ lệ hàng trả lại theo danh mục tháng đang chọn, %")
S4.add_chart(stacked_bar(CRR_T, ["F", "G"], [ACCENT, CTX], w=CM(10), h=8, fmt="0.0"), "B37")
title(S4, "L35", f'=FIXED({k_("canc")},0)&" đơn huỷ tháng "&Calc!$B$2', "Số đơn CANCELLED (xanh) và PENDING (xám) theo tháng đặt hàng")
S4.add_chart(line_chart(["canc", "pend"], w=CM(6), h=8, colors=(ACCENT, CTX), fmt="#,##0"), "L37")

# ================================================================== Hướng dẫn
SH.sheet_view.showGridLines = False
SH.column_dimensions["B"].width = 120
for i, t in enumerate([
    "Dashboard Excel tham chiếu — Sales Project 1 (Nexa Home)",
    "Mục đích: file ĐỐI CHIẾU SỐ cho mentor (và học viên sau buổi 1). Cùng bố cục và cùng số với Nexa_Sales_P1_Dashboard.html.",
    "Cách dùng: ở sheet Tong_quan chọn Tháng báo cáo, Kênh bán, Quốc gia trong 3 ô vàng; 4 sheet tự tính lại.",
    "Dựng bằng công thức thường (SUMIFS, COUNTIFS, INDEX/MATCH, LARGE) trên sheet Data — không dùng Power Pivot. "
    "Bài của học viên dựng bằng Power Query + Mô hình Dữ liệu + DAX theo hướng dẫn; số phải khớp file này.",
    "Sheet Data: tháng × cửa hàng (đã bỏ dòng trùng, chỉ đơn COMPLETED); Data_dm: × danh mục; Data_ld: × lý do trả hàng; "
    "Calc (ẩn): phép tính trung gian cho thẻ KPI, câu Điểm chính và biểu đồ.",
    "Quy ước: Δ = số Kế toán (finance_revenue.xlsx) − dữ liệu đơn; vượt ngưỡng khi |Δ| > "
    f"{THR_ABS:,.0f} USD và |Δ%| > {THR_PCT * 100:.0f}% ở mức tháng × cửa hàng; bút toán HQ-ADJ không thuộc cửa hàng nào.",
    "Kiểm: python _mentor/xuat_pdf_kiem_tra.py tính lại bằng LibreOffice, xuất PDF 4 sheet và đặt cạnh ảnh bản HTML.",
], 2):
    SH[f"B{i}"] = t
    SH[f"B{i}"].font = f(size=14 if i == 2 else 10, bold=i == 2)
    SH[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")
for ws in (SD, SM, SL):
    ws.sheet_state = "visible"
wb.calculation.fullCalcOnLoad = True
wb.active = 0
wb.save(OUT)

# openpyxl ghi <definedNames/> rỗng khi không có tên nào ngoài vùng in → xoá cho gọn (kiem_xlsx.py báo nếu còn)
import re
import shutil
import zipfile
tmp = OUT.with_suffix(".tmp")
with zipfile.ZipFile(OUT) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        data = zi.read(it.filename)
        if it.filename == "xl/workbook.xml":
            data = data.replace(b"<definedNames/>", b"")
        if it.filename.startswith("xl/charts/chart"):    # nhãn số dùng định dạng riêng; bỏ run rỗng
            data = re.sub(rb'<numFmt formatCode="([^"]*)"/>', rb'<numFmt formatCode="\1" sourceLinked="0"/>', data)
            data = data.replace(b"<a:r><a:t></a:t></a:r>", b"")
        zo.writestr(it, data)
shutil.move(tmp, OUT)
print("saved", OUT.name, "·", len(rows), "dòng Data")
