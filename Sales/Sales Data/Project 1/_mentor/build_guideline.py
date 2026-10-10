# -*- coding: utf-8 -*-
"""Dựng ../Dashboard_Guideline (dành cho mentor).docx

Khuôn: guideline của Bank Project 1 (giữ nguyên style, khổ giấy, bảng màu) — xoá nội dung cũ, viết nội dung Sales.
Mọi con số đọc từ ref_numbers.json, Instructor Key và chính file dataset => chạy lại sau generate_dataset.py.
Hình: lấy PNG xuất từ các file .drawio trong ../diagrams (sửa sơ đồ ở .drawio, chạy build_diagrams.py rồi chạy lại file này).
"""
import copy
import json
from pathlib import Path

import docx
import pandas as pd
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
P1 = HERE.parent
BANK = P1.parents[2] / "Bank" / "Project 1" / "Dashboard_Guideline (dành cho mentor).docx"
OUT = P1 / "Dashboard_Guideline (dành cho mentor).docx"
PNG = P1 / "diagrams" / "png"
R = json.loads((HERE / "ref_numbers.json").read_text(encoding="utf-8"))
KEY = pd.read_excel(HERE / "Nexa_Sales_P1_Instructor_Key.xlsx", sheet_name=None)
DS = pd.read_excel(P1 / "Nexa_Sales_P1_Dataset.xlsx", sheet_name=None)


def nf(v, d=0):
    s = f"{abs(v):,.{d}f}".replace(",", "§").replace(".", ",").replace("§", ".")
    return ("−" if v < 0 else "") + s


def pct(v, d=1):
    return nf(v * 100, d) + "%"


# ------------------------------------------------------------------ khuôn tài liệu
doc = docx.Document(BANK)
body = doc.element.body
for el in list(body):
    if el.tag != qn("w:sectPr"):
        body.remove(el)
STY = {s.name: s for s in doc.styles}          # tên style pandoc ("Heading 1") — tra trực tiếp, tránh ánh xạ của python-docx
num_part = doc.part.numbering_part.element
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
BULLET_ABS = None
DEC_ABS = None
for an in num_part.findall(W + "abstractNum"):
    fmt = an.find(W + "lvl").find(W + "numFmt").get(W + "val")
    if fmt == "bullet" and BULLET_ABS is None and an.find(W + "lvl").find(W + "lvlText").get(W + "val") == "•":
        BULLET_ABS = an.get(W + "abstractNumId")
    if fmt == "decimal" and DEC_ABS is None:
        DEC_ABS = an.get(W + "abstractNumId")
_next_num = [5000]


def new_num(abs_id):
    _next_num[0] += 1
    n = OxmlElement("w:num")
    n.set(qn("w:numId"), str(_next_num[0]))
    a = OxmlElement("w:abstractNumId")
    a.set(qn("w:val"), abs_id)
    n.append(a)
    ov = OxmlElement("w:lvlOverride")
    ov.set(qn("w:ilvl"), "0")
    so = OxmlElement("w:startOverride")
    so.set(qn("w:val"), "1")
    ov.append(so)
    n.append(ov)
    num_part.append(n)
    return _next_num[0]


def runs(p, text):
    """**đậm** và `code` đơn giản trong một đoạn."""
    import re
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        else:
            p.add_run(part)
    return p


def H1(t): doc.add_paragraph(t, style=STY["Heading 1"])
def H2(t): doc.add_paragraph(t, style=STY["Heading 2"])
def H3(t): doc.add_paragraph(t, style=STY["Heading 3"])
def P(t, first=False): return runs(doc.add_paragraph(style=STY["First Paragraph"] if first else STY["Body Text"]), t)


def LIST(items, numbered=False):
    nid = new_num(DEC_ABS if numbered else BULLET_ABS)
    for it in items:
        p = runs(doc.add_paragraph(style=STY["Compact"]), it)
        pPr = p._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0")
        ni = OxmlElement("w:numId"); ni.set(qn("w:val"), str(nid))
        numPr.append(il); numPr.append(ni)
        pPr.insert(0, numPr) if pPr.find(qn("w:pStyle")) is None else pPr.find(qn("w:pStyle")).addnext(numPr)


def CODE(text):
    p = doc.add_paragraph(style=STY["Source Code"])
    for i, line in enumerate(text.strip("\n").split("\n")):
        if i:
            p.add_run().add_break()
        p.add_run(line)


def TABLE(head, rows, widths=None, align=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(head))
    t.style = STY["Table"]
    tblPr = t._tbl.tblPr
    w_ = OxmlElement("w:tblW"); w_.set(qn("w:type"), "pct"); w_.set(qn("w:w"), "5000"); tblPr.append(w_)
    b = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        e = OxmlElement(f"w:{side}")
        for k, v in [("val", "single"), ("sz", "4"), ("space", "0"), ("color", "BED0B4")]:
            e.set(qn(f"w:{k}"), v)
        b.append(e)
    tblPr.append(b)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for ri, row in enumerate([head] + rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri].cells[ci]
            p = cell.paragraphs[0]
            p.style = STY["Compact"]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            if align and ri and align[ci] == "r":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if ri == 0:
                r = p.add_run(str(val)); r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); r.font.size = Pt(8.5)
                shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "447A2A")
                cell._tc.get_or_add_tcPr().append(shd)
            else:
                runs(p, str(val))
                for r in p.runs:
                    r.font.size = Pt(9)
    if widths:
        t.autofit = False
        grid = t._tbl.tblGrid
        for ci, (gc, wcm) in enumerate(zip(grid.findall(qn("w:gridCol")), widths)):
            gc.set(qn("w:w"), str(int(wcm * 567)))
            for row in t.rows:
                row.cells[ci].width = Cm(wcm)
    doc.add_paragraph(style=STY["Compact"])
    return t


def FIG(png, caption, width_cm=17.5):
    p = doc.add_paragraph(style=STY["Captioned Figure"])
    p.add_run().add_picture(str(PNG / png), width=Cm(width_cm))
    doc.add_paragraph(caption, style=STY["Image Caption"])


# ------------------------------------------------------------------ số liệu dùng nhiều lần
rows = R["rows"]
cur, lm, sp = R["month"]["current"], R["month"]["lm"], R["month"]["sply"]
ytd, ytd_sp = R["ytd"]["current"], R["ytd"]["sply"]
thr = R["threshold"]
br = R["bridge_2020"]
anom = KEY["planted_anomalies"]
over = KEY["over_threshold"]
over20 = over[over.month_id // 100 == 2020].copy()
adj_cols = [c for c in over20.columns if c.startswith("adj_")]
CAUSE_VI = {"adj_online_timing": "Online lệch thời điểm", "adj_refund_lag": "Hoàn tiền hạch toán trễ",
            "adj_cancelled_booked": "Đơn huỷ vẫn ghi nhận", "adj_missing_upload": "Finance thiếu dữ liệu ngày",
            "adj_vat_included": "Doanh thu gồm VAT", "adj_duplicate_posting": "Bút toán ghi trùng"}
over20["cause"] = over20[adj_cols].abs().idxmax(axis=1).map(CAUSE_VI)
fin_rows = rows["finance_revenue"]
hq20 = float(anom[(anom.type == "E_MANUAL_ADJ") & (pd.to_datetime(anom.date_from).dt.year == 2020)].effect_usd.sum())
n_canc_2020 = int(((anom.type == "A_CANCELLED_BOOKED") & (pd.to_datetime(anom.date_from).dt.year == 2020)).sum())
mstr = lambda m: f"{m % 100:02d}/{m // 100}"
PQ = json.loads((HERE / "pq_steps.json").read_text(encoding="utf-8"))
KIT = P1 / "excel_build"
DAX_TXT = (KIT / "measures.dax").read_text(encoding="utf-8")


def dax_sections(txt):
    """[(tiêu đề nhóm, [dòng code measure…])] theo các dòng // ===== trong measures.dax"""
    out, cur_ = [], None
    for line in txt.splitlines():
        if line.startswith("// ====="):
            cur_ = (line[8:].strip(), []); out.append(cur_)
        elif cur_ is not None and line.strip():
            cur_[1].append(line)
    return out


DAXS = dax_sections(DAX_TXT)
_fin = pd.read_excel(P1 / "finance_revenue.xlsx", sheet_name="finance_revenue")
fin_online = int(_fin.fin_store_code.isin(DS["map_store_finance"].query("store_key == 0").fin_store_code).sum())
MEAS = [l.split(":=")[0].strip() for _, ls in DAXS for l in ls if ":=" in l and not l.startswith(" ")]

# ------------------------------------------------------------------ NỘI DUNG
H1("Hướng dẫn xây dựng Dashboard Excel: Kết quả bán hàng và đối soát doanh thu (Nexa Home · Sales Project 1)")
P("Cập nhật: 10/10/2026 · Dataset v" + R["version"] + f" · seed {R['seed']}", first=True)

H2("Giới thiệu và cách dùng tài liệu")
P("Tài liệu này hướng dẫn dựng dashboard **Excel** 4 sheet (Tổng quan, Doanh thu, Đối soát, Hoàn trả & chất lượng) "
  "từ dataset Nexa Home, rồi dùng chính dashboard đó để đối soát doanh thu với file Finance và viết báo cáo. "
  "Quy trình gồm 4 tầng: **Power Query (M)** nạp và làm sạch → **Data Model** (Power Pivot) nối các bảng → "
  "**DAX** tính chỉ tiêu → **PivotTable, PivotChart, Slicer và CUBEVALUE** trình bày. "
  "Đây là project đầu tiên của nhóm Sales, mức hỗ trợ của mentor khoảng 80 %: mentor đi cùng học viên qua từng bước, "
  "học viên tự làm phần giải thích chênh lệch.", first=True)
P("**Đối tượng.** Mentor và trợ giảng dựng bản tham chiếu; học viên làm theo từng bước. **Bước 6** chứa số đối soát "
  "và đáp án các chênh lệch, phải ẩn khi dùng tài liệu này làm đề.")
P("**Sản phẩm đầu ra.** Một file .xlsx gồm các query Power Query, Data Model có quan hệ và bộ measure DAX, cùng 4 sheet "
  "dashboard; kèm báo cáo 8 phần theo guild_report_template.pdf, trong đó phần chính là bảng giải thích chênh lệch "
  "vượt ngưỡng năm 2020.")

H3("Yêu cầu phần mềm")
LIST(["**Excel cho Windows**, bản Microsoft 365 hoặc Excel 2016 trở lên. Excel cho Mac và Excel trên web không có "
      "Power Pivot nên không làm được Bước 3–4.",
      "**Bật add-in Power Pivot.** File → Options → Add-ins → ô Manage chọn **COM Add-ins** → Go… → tick "
      "**Microsoft Power Pivot for Excel** → OK. Ribbon có thêm tab **Power Pivot**.",
      "Power Query có sẵn trong tab Data (nhóm Get & Transform Data), không cần cài thêm.",
      "Toàn bộ code M và DAX nằm trong thư mục excel_build/; tài liệu này in lại đúng code đó."])

H3("Bộ file đầu vào")
TABLE(["File", "Công dụng", "Ai được xem"], [
    ["Nexa_Sales_P1_Dataset.xlsx", "Dữ liệu đơn hàng, hoàn trả, các bảng dim, bảng map mã Finance, định nghĩa KPI, giao dịch mẫu; mỗi sheet là một Excel Table tên tbl_<sheet>", "Học viên"],
    ["finance_revenue.xlsx", "Doanh thu Finance đã xác nhận theo ngày × mã cửa hàng Finance (bản rút gọn)", "Học viên"],
    ["excel_build/m/*.pq", f"{len(PQ['order'])} query M theo đúng thứ tự tạo, mỗi query một file; dán vào Advanced Editor", "Học viên"],
    ["excel_build/measures.dax", f"{len(MEAS)} measure DAX theo nhóm, dạng `Tên := công thức`", "Học viên"],
    ["Nexa_Sales_P1_Dashboard.html", "Bản dashboard tham chiếu (bố cục, màu, câu chữ) mở bằng trình duyệt", "Học viên (sau buổi 1) / mentor"],
    ["Nexa_Sales_P1_Dashboard.xlsx", "File đối chiếu số dựng bằng SUMIFS (không dùng Data Model): chọn tháng/kênh rồi so với dashboard của mình", "Học viên (sau buổi 1) / mentor"],
    ["Slide_Project_1.pptx", "Slide giảng 8 phần; phụ lục mentor đặt ẩn", "Học viên + mentor"],
    ["guild_report_template.pdf", "Khung báo cáo 8 phần", "Học viên"],
    ["diagrams/*.drawio (+ .pdf)", "Sơ đồ pipeline, mô hình, quy trình đối soát, cây KPI, pipeline Power Query; sửa ở .drawio rồi xuất lại", "Học viên + mentor"],
    ["_mentor/Nexa_Sales_P1_Instructor_Key.xlsx", "Đáp án: 31 chênh lệch cài sẵn, cầu nối từng tháng × cửa hàng", "Chỉ mentor"],
    ["_mentor/generate_dataset.py, build_excel_kit.py", "Sinh lại dữ liệu (seed cố định) và bộ M/DAX, kèm kiểm số dòng từng bước", "Chỉ mentor"],
], widths=[5.2, 9, 3.3])

H3("Dữ liệu trong một cái nhìn")
TABLE(["Nội dung", "Giá trị"], [
    ["Bối cảnh", "Nexa Home bán điện máy qua cửa hàng (POS) và website/app; Finance cần biết doanh thu thật và lý do lệch số"],
    ["Nguồn nền", "Global Electronics Retailer (Maven Analytics), giữ nguyên địa lý: 66 cửa hàng ở 8 quốc gia + 1 kênh online"],
    ["Khoảng thời gian", "Đơn hàng 01/01/2016 → 20/02/2021; ngày trích dữ liệu 20/02/2021"],
    ["Quy mô", f"{nf(rows['orders_total'])} đơn · {nf(rows['fact_sales_line'])} dòng bán (gồm 120 dòng trùng) · "
               f"{nf(rows['fact_return'])} lần trả hàng · {nf(fin_rows)} dòng Finance"],
    ["Tiền tệ", "USD, chưa gồm VAT; VAT giả định 10 % chỉ xuất hiện ở file Finance"],
    ["Tháng báo cáo mặc định", f"{mstr(R['report_month'])}; so với tháng trước (LM) và cùng tháng năm trước (SPLY)"],
    ["Phạm vi đối soát", "Năm 2020, mức tháng × cửa hàng × kênh"],
    ["Ngưỡng", f"Vượt ngưỡng khi |Δ| > {nf(thr['abs_usd'])} USD **và** |Δ %| > {pct(thr['pct'], 0)}"],
], widths=[4.5, 13])

H3("Lộ trình 7 bước")
TABLE(["Bước", "Việc chính", "Công cụ Excel", "Thời lượng gợi ý"], [
    ["1. Hiểu dữ liệu", "Đọc README, Data_Dictionary, KPI_Definitions, sample_transactions", "Mở file dataset", "30 phút"],
    ["2. Power Query (M)", f"{len(PQ['order'])} query theo pipeline: nạp, đặt kiểu, bỏ dòng trùng, map mã Finance → fact_finance", "Data → Get Data, Advanced Editor", "60 phút"],
    ["3. Data Model", "Quan hệ star schema, Mark as Date Table, quan hệ inactive cho delivery_date", "Power Pivot → Manage → Diagram View", "30 phút"],
    ["4. DAX", f"{len(MEAS)} measure trong bảng _Measures, gồm measure chuỗi cho tiêu đề", "Power Pivot → Measures → New Measure", "90 phút"],
    ["5. Dashboard", "4 sheet theo bảng visual → đối tượng Excel", "PivotTable, PivotChart, Slicer, CUBEVALUE", "120 phút"],
    ["6. Đối soát", "So số với Bước 6; lập bảng giải thích chênh lệch năm 2020", "Sheet Doi_soat", "60 phút"],
    ["7. Báo cáo", "Viết báo cáo 8 phần theo template", "—", "Ở nhà"],
], widths=[3.2, 7.3, 4.8, 2.2])

H3("Quy ước trong tài liệu")
LIST(["Tên bảng và cột viết dạng `fact_sales_line[order_date]`; tên measure viết dạng `[Net Revenue]`.",
      f"Mọi con số mẫu lấy từ chính dataset, tại tháng báo cáo {mstr(R['report_month'])} nếu không ghi khác.",
      "“LM” là tháng trước (Last Month), “SPLY” là cùng tháng năm trước (Same Period Last Year).",
      "Δ (chênh lệch) luôn tính **Finance − dữ liệu đơn**: Δ dương nghĩa là Finance ghi nhận cao hơn.",
      "Số viết theo kiểu Việt Nam: dấu chấm phân cách nghìn, dấu phẩy thập phân."])
FIG("01_pipeline_du_lieu.png", "Hình 1: Pipeline từ nguồn dữ liệu đến báo cáo (nguồn: diagrams/01_pipeline_du_lieu.drawio)")

# ---------------------------------------------------------------- Bước 1
H2("Bước 1. Hiểu dữ liệu và chỉ tiêu")
P("Dataset gồm 2 bảng fact, 7 bảng dim, 1 bảng map và file Finance riêng. Nắm grain (mỗi dòng là gì) của từng bảng "
  "trước khi viết bất kỳ measure nào.", first=True)
H3("1.1 Các bảng trong dataset")
cnt = {k: len(v) for k, v in DS.items()}
TABLE(["Bảng", "Vai trò", "Grain (mỗi dòng là)", "Khoá", "Số dòng", "Nhóm"], [
    ["fact_sales_line", "Fact", "1 sản phẩm trong 1 đơn", "order_number + line_item", nf(cnt["fact_sales_line"]), "Bắt buộc"],
    ["fact_return", "Fact", "1 lần trả hàng của 1 dòng", "return_id", nf(cnt["fact_return"]), "Bắt buộc"],
    ["finance_revenue (file riêng)", "Fact", "1 ngày × 1 mã cửa hàng Finance", "—", nf(fin_rows), "Bắt buộc"],
    ["dim_date", "Dim", "1 ngày", "date", nf(cnt["dim_date"]), "Bắt buộc"],
    ["dim_store", "Dim", "1 cửa hàng (0 = online)", "store_key", nf(cnt["dim_store"]), "Bắt buộc"],
    ["dim_product", "Dim", "1 sản phẩm", "product_key", nf(cnt["dim_product"]), "Bắt buộc"],
    ["dim_channel", "Dim", "1 kênh", "channel_code", nf(cnt["dim_channel"]), "Bắt buộc"],
    ["dim_promotion", "Dim", "1 chương trình KM", "promo_code", nf(cnt["dim_promotion"]), "Bắt buộc"],
    ["dim_order_status", "Dim", "1 trạng thái đơn", "order_status", nf(cnt["dim_order_status"]), "Bắt buộc"],
    ["map_store_finance", "Map", "1 mã Finance trong 1 khoảng hiệu lực", "fin_store_code", nf(cnt["map_store_finance"]), "Bắt buộc"],
    ["dim_customer", "Dim", "1 khách hàng", "customer_key", nf(cnt["dim_customer"]), "Bổ sung"],
    ["fx_rate", "Tham khảo", "1 ngày × 1 tiền tệ", "date + currency_code", nf(cnt["fx_rate"]), "Bổ sung"],
], widths=[3.6, 1.6, 4.2, 3.6, 1.8, 1.8], align="lllllr")
P("Nhóm “Bắt buộc” đủ để dựng 4 trang và đối soát. Nhóm “Bổ sung” dành cho phần sáng tạo (20 điểm trong rubric), "
  "ví dụ phân tích theo vùng khách hàng hoặc quy đổi tiền tệ.")
H3("1.2 Từ chỉ tiêu nghiệp vụ đến cột dữ liệu")
kpi = DS["KPI_Definitions"]
TABLE(["Chỉ tiêu", "Công thức", "Ghi chú"], [[r.kpi, r.formula, "" if pd.isna(r.scope_note) else r.scope_note]
                                            for r in kpi.itertuples()], widths=[3.6, 6.4, 7.5])
FIG("04_cay_kpi.png", "Hình 2: Cây KPI (nguồn: diagrams/04_cay_kpi.drawio)")
H3("1.3 Chín quy tắc dữ liệu cần nhớ")
LIST([
    "**Chỉ đơn COMPLETED.** CANCELLED (khách huỷ) và PENDING (online chưa giao tại 20/02/2021) không phải doanh thu. "
    f"Có {nf(rows['orders_cancelled'])} đơn huỷ và {nf(rows['orders_pending'])} đơn đang giao.",
    "**Dòng trùng.** 120 dòng xuất hiện 2 lần y hệt; khoá dòng là order_number + line_item. Bỏ trùng trước khi tính.",
    "**Khuyến mãi áp cho cả đơn.** discount_amount đã tính sẵn ở từng dòng = quantity × unit_price × discount_pct.",
    "**Hoàn trả tính theo ngày trả.** Returns lấy fact_return[return_date], không phải ngày đặt của đơn gốc; một dòng có thể trả một phần.",
    "**Online có 2 mốc ngày.** Dữ liệu đơn dùng order_date; Finance ghi nhận online theo delivery_date (trung bình 4–5 ngày sau).",
    "**POS không có delivery_date.** Khách nhận hàng ngay tại quầy, Finance ghi nhận đúng ngày bán.",
    "**Mã Finance khác store_key.** Dùng map_store_finance; cửa hàng 38 (Anh) đổi mã GBR-03 → GBR-03N từ 01/07/2019, "
    "phải map theo valid_from/valid_to. Mã HQ-ADJ không thuộc cửa hàng nào.",
    "**Finance hạch toán hoàn tiền trễ.** refunds trong file Finance có thể lệch 0–3 ngày so với return_date.",
    "**USD, chưa VAT.** Mọi số tiền phía đơn hàng là USD chưa VAT; so với cột net_revenue của Finance, không so với total_incl_vat.",
])
H3("1.4 Cột được phép để trống")
P("fact_sales_line[delivery_date] trống với mọi đơn POS và mọi đơn online chưa COMPLETED; dim_store[square_meters] trống "
  "ở dòng online; dim_promotion[window_from_mmdd]/[window_to_mmdd] trống với NONE và WEB10; map_store_finance[store_key] "
  "trống ở HQ-ADJ. Mọi cột khác phải đầy đủ.", first=True)

# ---------------------------------------------------------------- Bước 2
H2("Bước 2. Power Query (M): nạp và làm sạch theo pipeline")
P("Dữ liệu ở đây **chưa sạch hoàn toàn**: có dòng trùng và file Finance dùng mã cửa hàng riêng. Hai việc làm sạch này "
  "là một phần của bài. Trong Power Query, mỗi thao tác chuột được ghi thành một **Applied Step**; trong code M, mỗi "
  "step là một dòng `TênBước = Hàm(bước trước, …)` nằm giữa `let` và `in`. Đọc từ trên xuống là đọc đúng thứ tự xử lý.",
  first=True)
FIG("05_power_query_pipeline.png", "Hình 3: Pipeline Power Query — mỗi ô là một Applied Step, số là số dòng sau bước "
    "(nguồn: diagrams/05_power_query_pipeline.drawio)")
P("Đọc Hình 3 theo làn: ô cam là bước **thay đổi số dòng hoặc ghép bảng** (bỏ trùng, merge, lọc hiệu lực) — chỗ hay sai "
  "nhất, phải kiểm số dòng ngay sau bước đó. Mũi tên nét đứt là phụ thuộc giữa các query: fact_return cần "
  "fact_sales_line; fact_finance cần finance_raw và map_store_finance.")

H3("2.1 Chuẩn bị: thư mục dữ liệu và tham số SourceFolder")
LIST(["Chép Nexa_Sales_P1_Dataset.xlsx và finance_revenue.xlsx vào cùng một thư mục, ví dụ `C:\\NexaP1\\`.",
      "Mở một workbook Excel trắng, lưu ngay thành `Nexa_P1_<HoTen>.xlsx` (file nộp).",
      "Data → Get Data → Launch Power Query Editor → Home → Manage Parameters → New Parameter: Name `SourceFolder`, "
      "Type **Text**, Current Value là đường dẫn thư mục **kết thúc bằng dấu \\\\**.",
      "Mọi query nguồn đọc file qua `SourceFolder & \"tên file\"`; đổi máy chỉ cần sửa tham số này rồi Refresh All."],
     numbered=True)
P("Có hai cách tạo từng query, kết quả giống nhau:", first=True)
LIST(["**Thao tác chuột** (khuyên dùng cho fact_sales_line và fact_finance để hiểu bước): Get Data → From File → From "
      "Excel Workbook → chọn file → trong Navigator chọn **Table** `tbl_…` (biểu tượng bảng xanh), không chọn sheet cùng tên → "
      "Transform Data; làm từng bước và so tên bước ở khung Applied Steps với bảng bên dưới.",
      "**Dán code**: Home → New Source → Other Sources → Blank Query → Advanced Editor → dán nội dung file "
      "excel_build/m/<số>_<tên>.pq → Done → đổi tên query (khung Query Settings → Name) đúng như tên file, bỏ số thứ tự."])

H3("2.2 Thứ tự tạo query")
STEPS = PQ["steps"]
TAIL = {"SourceFolder": "Tham số đường dẫn", "_Measures": "Bảng rỗng chứa measure"}
q_rows = []
for i, q in enumerate(PQ["order"], 1):
    if q in STEPS:
        main = " → ".join(st[0] for st in STEPS[q][1:])
        n = nf(STEPS[q][-1][2])
    elif q in PQ["dims"]:
        main, n = "Navigation → Changed Type", nf(PQ["dims"][q])
    else:
        main, n = TAIL[q], "—"
    q_rows.append([str(i), q, main, PQ["load"].get(q, "Connection + Data Model"), n])
TABLE(["#", "Query", "Các bước (sau Source)", "Nạp (Close & Load To)", "Số dòng"], q_rows,
      widths=[0.9, 3.4, 7.4, 3.9, 1.9], align="llllr")
P("Tạo theo đúng thứ tự này: query sau tham chiếu query trước (fact_return dùng fact_sales_line; fact_finance dùng "
  "finance_raw và map_store_finance), tạo ngược sẽ báo lỗi *Expression.Error: The name … wasn't recognized*.".replace("*", ""))

STEP_VI = {
    "Source": "Mở file từ thư mục SourceFolder (Excel.Workbook trả về danh sách sheet và Table trong file)",
    "Navigation": "Chọn đúng Excel Table tbl_… — có sẵn tiêu đề cột, không dính dòng trống",
    "Changed Type": "Đặt kiểu từng cột: Date, Whole number, Fixed decimal (Currency), Text",
    "Removed Duplicates": "Home → Remove Rows → Remove Duplicates trên toàn bộ cột: bỏ 120 dòng trùng hoàn toàn",
    "Added gross_amount": "Add Column → Custom Column: quantity × unit_price, kiểu Fixed decimal",
    "Added net_amount": "Custom Column: gross_amount − discount_amount",
    "Merged Queries": "Home → Merge Queries (Left Outer)",
    "Expanded sales": "Bấm nút ⇄ ở cột sales, chỉ tick store_key, channel_code, product_key, bỏ tick “Use original column name as prefix”",
    "Changed Type2": "Đặt kiểu cho 3 cột vừa lấy",
    "Expanded map": "Bấm nút ⇄ ở cột map, tick store_key, valid_from, valid_to",
    "Filtered Valid": "Lọc tạm một điều kiện bất kỳ trên posting_date rồi sửa công thức trên thanh công thức thành `each [posting_date] >= [valid_from] and [posting_date] <= [valid_to]`",
    "Removed Columns": "Bỏ valid_from, valid_to (chỉ dùng để lọc)",
}
MERGE_VI = {"fact_return": "Ghép dòng bán gốc theo **order_number + line_item** (giữ Ctrl để chọn 2 cột ở cả hai bảng)",
            "fact_finance": "Ghép map_store_finance theo fin_store_code"}


def pq_block(q, sub, note=None):
    H3(sub)
    if note:
        P(note, first=True)
    rows_ = []
    for k, (stp, fn, n) in enumerate(STEPS[q], 1):
        why = MERGE_VI[q] if stp == "Merged Queries" else (
            "Home → Reference từ query finance_raw" if fn.startswith("=") else STEP_VI.get(stp, ""))
        rows_.append([str(k), stp, fn, "—" if n is None else nf(n), why])
    TABLE(["#", "Applied Step", "Hàm M", "Số dòng", "Thao tác / ý nghĩa"], rows_,
          widths=[0.7, 3.1, 4.3, 1.5, 7.9], align="lllrl")
    CODE(PQ["queries"][q]["code"])


pq_block("fact_sales_line", "2.3 Pipeline fact_sales_line: bỏ dòng trùng, thêm cột tiền",
         "Bảng lớn nhất và là nơi duy nhất có dòng trùng. Bỏ trùng **sau** khi đặt kiểu, để hai dòng giống nhau về giá trị "
         "nhưng khác kiểu (ví dụ ngày dạng text) không bị coi là khác nhau.")
P(f"Kiểm tra: thanh trạng thái của Power Query Editor chỉ hiện “999+ ROWS” và Column profile mặc định chỉ xét 1.000 dòng đầu; muốn đếm đủ, chọn Transform → "
  f"Count Rows ở một bản Reference tạm (xoá sau khi xem). Sau Removed Duplicates phải còn {nf(rows['fact_sales_line_clean'])} dòng.")
pq_block("fact_return", "2.4 Pipeline fact_return: ghép cửa hàng, kênh, sản phẩm từ dòng bán gốc",
         "fact_return không có store_key nên không lọc được theo cửa hàng, kênh hay danh mục. Merge với fact_sales_line "
         "(đã bỏ trùng, nên mỗi dòng trả ghép đúng một dòng bán) để mang 3 khoá này sang.")
P(f"Kiểm tra: số dòng trước và sau merge đều là {nf(rows['fact_return'])}; nếu tăng là do merge với fact_sales_line "
  "**chưa** bỏ trùng. Lọc thử store_key = null: phải không còn dòng nào.")
H3("2.5 Pipeline Finance: finance_raw, map_store_finance → fact_finance")
P("finance_raw và map_store_finance là query trung gian (chỉ tạo connection, không nạp vào Data Model). Cả hai đều gồm "
  "Source → Navigation → Changed Type như các bảng dim. File Finance chỉ có mã cửa hàng của Finance (fin_store_code); "
  "fact_finance gắn store_key bằng bảng map theo **mã và ngày hiệu lực**:", first=True)
pq_block("fact_finance", "Các bước của fact_finance")
P(f"Kiểm tra: fact_finance có đúng {nf(fin_rows)} dòng như finance_raw, và đúng 3 dòng HQ-ADJ có store_key = null "
  "(giữ lại để giải thích ở tổng chuỗi, không gán cho cửa hàng nào). Với dữ liệu hiện tại bước Filtered Valid giữ nguyên "
  "số dòng: mã cũ GBR-03 chỉ có bút toán đến 27/06/2019, mã mới GBR-03N từ 01/07/2019. Bước này là chốt an toàn — nếu "
  "Finance còn hạch toán mã cũ sau ngày đổi, dòng đó bị loại và số dòng giảm, là tín hiệu cần hỏi lại Finance.")
H3("2.6 Các bảng dim và bảng _Measures")
dim_note = {"dim_date": "1 dòng/ngày 2016–2021; month_label dạng 12/2020", "dim_store": "66 cửa hàng + dòng 0 = Online",
            "dim_product": "category, subcategory, brand", "dim_channel": "POS, ONLINE",
            "dim_promotion": "Mã khuyến mãi", "dim_order_status": "COMPLETED, CANCELLED, PENDING"}
TABLE(["Query", "Excel Table nguồn", "Số dòng", "Ghi chú"],
      [[d, "tbl_" + d, nf(n), dim_note[d]] for d, n in PQ["dims"].items()], widths=[3.5, 4, 2, 8], align="llrl")
P("Mọi bảng dim cùng một mẫu 3 bước (Source → Navigation → Changed Type). Mẫu dim_date:")
CODE(PQ["queries"]["dim_date"]["code"])
P("Bảng `_Measures` là một bảng một dòng tạo bằng M, chỉ để chứa measure (Excel bắt buộc measure phải thuộc một bảng):")
CODE(PQ["queries"]["_Measures"]["code"])
H3("2.7 Nạp vào Data Model")
LIST(["Home → **Close & Load To…** (không bấm Close & Load, vì lệnh đó đổ mọi bảng ra sheet).",
      "Chọn **Only Create Connection** và tick **Add this data to the Data Model** cho các dim, fact_sales_line, "
      "fact_return, fact_finance và _Measures.",
      "finance_raw và map_store_finance: chỉ **Only Create Connection**, không tick Data Model.",
      "Lỡ nạp sai: Data → Queries & Connections → chuột phải query → Load To… để đổi lại.",
      "Kiểm số dòng: Power Pivot → Manage, mỗi tab bảng ghi số dòng ở góc dưới (Record: 1 of …)."], numbered=True)
H3("2.8 Số dòng cần khớp sau khi nạp")
TABLE(["Query", "Số dòng"], [
    ["fact_sales_line (trước / sau Removed Duplicates)", f"{nf(rows['fact_sales_line'])} → {nf(rows['fact_sales_line_clean'])}"],
    ["fact_return", nf(rows["fact_return"])],
    ["fact_finance (trong đó HQ-ADJ có store_key trống)", f"{nf(fin_rows)} (3)"],
    ["dim_date", nf(rows["dim_date"])], ["dim_store", nf(rows["dim_store"])], ["dim_product", nf(rows["dim_product"])],
    ["map_store_finance (connection)", nf(rows["map_store_finance"])],
], widths=[12, 5.5], align="lr")
P("Sai khác số dòng thường do quên bỏ trùng, chọn nhầm sheet thay vì Table, hoặc merge fact_return với bản "
  "fact_sales_line chưa bỏ trùng (fact_return dư dòng).")

# ---------------------------------------------------------------- Bước 3
H2("Bước 3. Data Model trong Power Pivot")
P("Mô hình hình sao: fact ở giữa, dim xung quanh, quan hệ 1 → nhiều và lọc một chiều. fact_sales_line và fact_finance "
  "không nối trực tiếp với nhau; chúng gặp nhau qua dim_date và dim_store. Mở Power Pivot → **Manage** → Home → "
  "**Diagram View** để thấy mọi bảng đã nạp.", first=True)
FIG("02_mo_hinh_du_lieu.png", "Hình 4: Mô hình dữ liệu (nguồn: diagrams/02_mo_hinh_du_lieu.drawio)")
H3("3.1 Quan hệ cần tạo")
TABLE(["Từ (1)", "Đến (nhiều)", "Trạng thái"], [
    ["dim_date[date]", "fact_sales_line[order_date]", "Active"],
    ["dim_date[date]", "fact_sales_line[delivery_date]", "Inactive (dùng USERELATIONSHIP)"],
    ["dim_date[date]", "fact_return[return_date]", "Active"],
    ["dim_date[date]", "fact_finance[posting_date]", "Active"],
    ["dim_store[store_key]", "fact_sales_line[store_key]", "Active"],
    ["dim_store[store_key]", "fact_return[store_key]", "Active (cột lấy ở Bước 2.4)"],
    ["dim_store[store_key]", "fact_finance[store_key]", "Active"],
    ["dim_product[product_key]", "fact_sales_line[product_key]", "Active"],
    ["dim_product[product_key]", "fact_return[product_key]", "Active (cột lấy ở Bước 2.4)"],
    ["dim_promotion[promo_code]", "fact_sales_line[promo_code]", "Active"],
    ["dim_order_status[order_status]", "fact_sales_line[order_status]", "Active"],
    ["dim_channel[channel_code]", "dim_store[channel_code]", "Active (kênh lọc qua cửa hàng)"],
], widths=[5.5, 6.5, 5.5])
LIST(["**Tạo quan hệ:** trong Diagram View kéo cột khoá ở bảng dim thả lên cột tương ứng ở bảng fact; hoặc Design → "
      "Create Relationship và chọn 2 bảng, 2 cột.",
      "**Quan hệ inactive:** kéo dim_date[date] lên fact_sales_line[delivery_date] lần thứ hai. Vì đã có quan hệ active "
      "giữa hai bảng, Excel tự để quan hệ mới ở trạng thái inactive (đường nét đứt). Kiểm ở Design → Manage Relationships: "
      "cột Active = No.",
      "**Không tạo** quan hệ giữa các bảng fact, và không nối map_store_finance (bảng này chỉ dùng trong Power Query)."])
H3("3.2 Cấu hình bắt buộc")
LIST(["**Mark as Date Table.** Chọn bảng dim_date → Design → Mark as Date Table → cột date. Bắt buộc để DATEADD và "
      "SAMEPERIODLASTYEAR chạy đúng khi lọc theo năm/tháng.",
      "**Sort by Column.** Ở Data View chọn cột dim_date[month_label] → Home → Sort by Column → sort theo month_id; nếu không, "
      "01/2017 đứng trước 12/2016.",
      "**Hide from Client Tools.** Chuột phải các khoá ở bảng fact (store_key, product_key, order_status, promo_code, các cột "
      "ngày) và cột ghi_chu của _Measures → Hide from Client Tools, để học viên luôn kéo trường từ bảng dim.",
      "**Định dạng.** Mỗi measure chọn định dạng ngay khi tạo (Currency 0 chữ số lẻ, Percentage 1–2 chữ số lẻ); "
      "PivotTable và CUBEVALUE dùng định dạng đó."])
H3("3.3 Ba phép thử mô hình (trước khi viết measure)")
P("Insert → PivotTable → **From Data Model** → New Worksheet (đặt tên sheet `_kiem_tra`). Kéo cột số vào Values, Excel "
  "tự tạo measure ngầm (Sum of …); với order_number đổi Value Field Settings → **Distinct Count**. Ra đúng số thì mô hình đã đúng:",
  first=True)
TABLE(["Phép thử", "Rows / Filters", "Values", "Phải ra"], [
    ["Đơn COMPLETED năm 2020", "Filters: dim_order_status[order_status] = COMPLETED; Rows: dim_date[year]",
     "fact_sales_line[order_number] · Distinct Count", f"2020: {nf(ytd['orders'])}"],
    ["Finance năm 2020 (gồm HQ-ADJ)", "Rows: dim_date[year]", "fact_finance[net_revenue] · Sum",
     f"2020: {nf(R['finance_year']['2020'], 2)}"],
    ["Finance gắn đúng kênh", "Rows: dim_channel[channel_code]", "fact_finance[fin_store_code] · Count",
     f"ONLINE: {nf(fin_online)} · POS: {nf(fin_rows - fin_online - 3)} · (blank): 3"],
], widths=[3.6, 6.1, 4.3, 3.5])
P("Phép thử thứ ba ra một dòng (blank) gồm 3 bút toán HQ-ADJ — đúng thiết kế. Nếu (blank) lớn hơn 3, fact_finance chưa gắn "
  "được store_key (xem lại 2.5).")

# ---------------------------------------------------------------- Bước 4
H2("Bước 4. Viết measure DAX trong Power Pivot")
P("Viết theo đúng thứ tự cây KPI: doanh thu → đơn hàng → Finance → chênh lệch → so sánh kỳ. Mỗi measure mới kiểm ngay "
  "bằng PivotTable ở sheet `_kiem_tra` trước khi dùng vào measure khác.", first=True)
H3("4.1 Tạo measure trong Excel")
LIST(["**Cách 1 — hộp thoại:** Power Pivot → Measures → **New Measure…** → Table name: `_Measures` → Measure name → "
      "Formula (phần sau dấu `=`) → Check formula → Category/Format → OK.",
      "**Cách 2 — Calculation Area:** Power Pivot → Manage → Data View → mở tab bảng _Measures → bấm một ô trong vùng lưới "
      "dưới bảng → gõ hoặc dán nguyên dòng `Tên := công thức` từ excel_build/measures.dax → Enter. Nhanh khi dán nhiều measure.",
      "Tên measure phải giữ **đúng như trong measures.dax**, vì các ô CUBEVALUE ở Bước 5 gọi measure theo tên.",
      "Sửa measure: Power Pivot → Measures → Manage Measures → Edit."], numbered=True)
H3("4.2 Nguyên tắc chung")
LIST(["Dùng DIVIDE thay cho phép chia “/” để tránh lỗi chia cho 0 (DIVIDE trả về trống).",
      "Bộ lọc trạng thái đặt **trong** measure (KEEPFILTERS), không đặt ở filter của PivotTable, để mọi sheet dùng chung một định nghĩa.",
      "Định dạng: tiền `#,##0` USD; tỷ lệ `0.0%`; chênh lệch `+#,##0;-#,##0;0`.",
      "Excel không có display folder như công cụ BI chuyên dụng: đặt tên measure có tiền tố nhóm thống nhất và tạo theo "
      "đúng thứ tự trong measures.dax để danh sách field dễ tìm.",
      "Excel không có cửa sổ chạy truy vấn DAX: kiểm measure bằng PivotTable (Bước 3.3) hoặc một ô CUBEVALUE."])
DAX_SUB = {1: "4.3 Nhóm Doanh thu", 2: "4.4 Nhóm Đơn hàng và giá", 3: "4.5 Nhóm Finance và đối soát",
           4: "4.6 Nhóm So sánh kỳ", 5: "4.7 Nhóm hỗ trợ biểu đồ Excel", 6: "4.8 Nhóm chuỗi cho tiêu đề hành động và câu Điểm chính"}
DAX_NOTE = {
    3: f"[Vượt ngưỡng] so một ô tháng × cửa hàng với ngưỡng {nf(thr['abs_usd'])} USD và {pct(thr['pct'], 0)}. "
       "[Số cửa hàng vượt ngưỡng] duyệt từng cặp tháng × cửa hàng (CROSSJOIN), nên đúng cả khi chọn một tháng lẫn cả năm. "
       "[Net Sales (ngày giao)] bật quan hệ inactive delivery_date để tính online theo ngày giao như Finance.",
    4: "Chỉ chạy đúng khi dim_date đã Mark as Date Table và slicer chọn **một** năm + **một** tháng.",
    5: "Excel không tô màu từng cột theo điều kiện trong PivotChart. Cách bền: tách một measure thành 2 chuỗi (mục nhấn / "
       "phần còn lại), vẽ cột chồng (stacked) rồi tô một chuỗi màu nhấn, một chuỗi xám. Đổi slicer, màu tự đi theo dữ liệu.",
    6: "Measure trả về chuỗi văn bản, đưa ra ô bằng CUBEVALUE (Bước 5) để tiêu đề và câu Điểm chính tự đổi theo slicer.",
}
for i, (title, lines) in enumerate(DAXS, 1):
    H3(DAX_SUB[i])
    if i in DAX_NOTE:
        P(DAX_NOTE[i], first=True)
    CODE("\n".join(lines))
H3("4.9 Năm lỗi hay gặp")
LIST(["**Quên lọc COMPLETED** → Net Sales cao hơn, AOV sai; Δ âm ở mọi cửa hàng.",
      "**Đếm Orders bằng COUNTROWS** → đếm dòng chứ không đếm đơn, AOV thấp bất thường (~1/2,4 giá trị đúng).",
      "**Trừ Returns theo ngày đặt của đơn gốc** → Returns không khớp tháng với Finance.",
      "**So với total_incl_vat** → Δ ≈ 10 % ở mọi ô.",
      "**Chưa Mark as Date Table** → các measure LM/SPLY trả về trống hoặc bằng số tháng hiện tại."])

# ---------------------------------------------------------------- Bước 5
H2("Bước 5. Dựng 4 sheet dashboard")
H3("5.1 Nguyên tắc trình bày: Storytelling with Data")
P("Dashboard dùng để ra quyết định trước cuộc họp chốt số, nên áp dụng 6 bài học của Cole Nussbaumer Knaflic "
  "(Storytelling with Data). Bản HTML tham chiếu làm theo các quy tắc dưới đây; khi chấm phần trải nghiệm (6 điểm) và "
  "lập luận, đối chiếu với bảng này.", first=True)
TABLE(["Bài học", "Áp dụng trong dashboard", "Lỗi hay gặp cần tránh"], [
    ["1. Hiểu bối cảnh", "Mỗi sheet mở bằng câu “Điểm chính” (Big Idea) tự sinh theo số: chuyện gì xảy ra, vì sao, cần làm gì.", "Sheet chỉ có biểu đồ, người xem phải tự đoán kết luận"],
    ["2. Chọn hình phù hợp", "Đường cho xu hướng; cột ngang xếp hạng cho danh mục/quốc gia; waterfall cho cầu doanh thu và phân rã Δ; bảng cho chi tiết.", "Pie/donut/treemap để so sánh, biểu đồ hai trục y, 3D"],
    ["3. Gỡ rối", "Tắt gridlines của sheet, bỏ chú giải khi chỉ có một chuỗi, ẩn field button của PivotChart, nhãn trực tiếp ở cuối thanh.", "Khung đậm, nền màu ở mọi tiêu đề, số trên mọi điểm"],
    ["4. Dẫn mắt", "Xám cho bối cảnh, một màu nhấn (xanh dương) cho tháng chọn / mục lớn nhất; đỏ chỉ cho ⚠ vượt ngưỡng; ▲▼ luôn kèm chữ.", "Mỗi chuỗi một màu, đỏ/xanh lá cho mọi thay đổi"],
    ["5. Nghĩ như nhà thiết kế", "Lưới cột thẳng hàng, định dạng số thống nhất, bảng màu an toàn cho người mù màu, font Segoe UI.", "Số lẻ khác nhau giữa các visual, chữ bị cắt"],
    ["6. Kể chuyện", "Tiêu đề hành động (“Máy tính chiếm 43% Net Sales”), mạch sheet: Tổng quan → Doanh thu đến từ đâu → Đối soát → Hoàn trả; chân trang ghi nguồn, as-of, quy ước Δ.", "Tiêu đề chỉ mô tả (“Cơ cấu theo danh mục”)"],
], widths=[3.2, 8.8, 5.5])
P("Trong Excel, tiêu đề hành động là **một ô công thức** đặt ngay trên biểu đồ (xoá Chart Title của biểu đồ), lấy measure "
  "chuỗi bằng CUBEVALUE, ví dụ `=CUBEVALUE(\"ThisWorkbookDataModel\",\"[Measures].[Tiêu đề danh mục]\",Slicer_year,"
  "Slicer_month,Slicer_channel_name,Slicer_country)`.")
H3("5.2 Khung chung: 4 sheet, 1 sheet pivot, 4 slicer")
LIST(["**Sheet.** `Tong_quan`, `Doanh_thu`, `Doi_soat`, `Hoan_tra` cho người xem; `_pivot` chứa mọi PivotTable nguồn của "
      "biểu đồ và bảng phụ (ẩn khi nộp: chuột phải tab → Hide).",
      "**Lưới.** Mỗi sheet dashboard: View → bỏ tick Gridlines và Headings; cột A rộng 2; cột B–Q rộng 11 (≈ 1.280 px); "
      "nền trắng, font Segoe UI.",
      "**4 slicer từ Data Model** (Insert → Slicer → tab Data Model): dim_date[year], dim_date[month], "
      "dim_channel[channel_name], dim_store[country]. Đặt ở Tong_quan, vùng B4:Q6; tab Slicer → Columns để slicer nằm ngang (Tháng 12 cột, Năm 6 cột), "
      "Slicer Settings → tick **Hide items with no data**.",
      "**Tên slicer.** Slicer Settings → *Name to use in formulas*: kiểm đúng `Slicer_year`, `Slicer_month`, "
      "`Slicer_channel_name`, `Slicer_country` (Excel thêm số 1 nếu tạo lặp — sửa lại cho khớp công thức).".replace("*", ""),
      "**Dùng chung slicer cho 4 sheet.** Copy slicer ở Tong_quan → Paste sang sheet khác: bản dán dùng chung bộ lọc "
      "(cùng slicer cache), chọn ở sheet nào cũng đổi cả 4 sheet.",
      "**Report Connections.** Chuột phải từng slicer → Report Connections → tick mọi PivotTable trên `_pivot`, **trừ** các "
      "pivot xu hướng 24 tháng và theo tháng với slicer Năm và Tháng (mục 5.9).",
      "**Màu.** Xám #B9B7B0 cho bối cảnh; xanh dương #2A78D6 nhấn; tốt = #1C5CAB, xấu = cam #EB6834; đỏ #D03B3B chỉ cho ⚠ vượt ngưỡng.",
      "Mở Nexa_Sales_P1_Dashboard.html để xem bố cục, màu và câu chữ mẫu; mở Nexa_Sales_P1_Dashboard.xlsx để so số."])
H3("5.3 Thẻ KPI và câu Điểm chính bằng CUBEVALUE")
P("CUBEVALUE lấy một measure từ Data Model theo đúng lựa chọn hiện tại của các slicer truyền vào. Mỗi thẻ KPI là 3 ô dọc: "
  "nhãn (xám, 9 pt), giá trị (đen, 20 pt đậm) và dòng so sánh (9 pt):", first=True)
CODE('''Nhãn       : Net Revenue (USD)
Giá trị    : =CUBEVALUE("ThisWorkbookDataModel","[Measures].[Net Revenue]",
                        Slicer_year,Slicer_month,Slicer_channel_name,Slicer_country)
So với LM  : =IFERROR(TEXT(CUBEVALUE("ThisWorkbookDataModel","[Measures].[Net Revenue % LM]",
                        Slicer_year,Slicer_month,Slicer_channel_name,Slicer_country),
                        "▲ 0.0%;▼ 0.0%")&" so với tháng trước","–")
Điểm chính : =CUBEVALUE("ThisWorkbookDataModel","[Measures].[Điểm chính Tổng quan]",
                        Slicer_year,Slicer_month,Slicer_channel_name,Slicer_country)
             &" "&CUBEVALUE("ThisWorkbookDataModel","[Measures].[Điểm chính Động lực]",
                        Slicer_year,Slicer_month,Slicer_channel_name,Slicer_country)''')
LIST(["Màu dòng so sánh: Home → Conditional Formatting → New Rule → Use a formula, ví dụ `=LEFT(C14,1)=\"▲\"` → chữ "
      "#1C5CAB; `=LEFT(C14,1)=\"▼\"` → chữ #EB6834 (với Return Rate và Δ thì đảo màu: tăng là xấu).",
      "Dải Điểm chính: gộp ô B8:Q9, Wrap Text, viền trái dày màu nhấn, nền #F2F6FC.",
      "Ô đọc tháng đang chọn (dùng ở dòng phụ tiêu đề): `=CUBEVALUE(\"ThisWorkbookDataModel\",\"[Measures].[Tháng chọn]\",Slicer_year,Slicer_month)`.",
      "Khi người xem chọn nhiều tháng, measure so sánh kỳ trả về trống → ô hiện “–” nhờ IFERROR, không hiện lỗi."])
H3("5.4 Biểu đồ: PivotChart, chart thường và bảng")
TABLE(["Loại hình", "Cách dựng trong Excel", "Ghi chú"], [
    ["Đường 24 tháng có điểm nhấn", "PivotTable trên _pivot: Rows dim_date[month_label], Filter dim_date[year] = 2019, 2020, "
     "Values measure; bảng phụ 3 cột: tháng, giá trị (=ô pivot), nhấn = IF(tháng = ô Tháng chọn, giá trị, NA()); "
     "Insert → Line chart thường từ bảng phụ", "Chuỗi giá trị màu xám; chuỗi nhấn chỉ hiện marker xanh. Pivot không nối slicer Năm/Tháng"],
    ["Waterfall (cầu doanh thu, phân rã Δ)", "Các ô CUBEVALUE: Gross Sales, −Discounts, −Returns, Net Revenue → Insert → "
     "Waterfall; bấm cột cuối → Set as total", "Excel 2016+; không làm được từ PivotTable"],
    ["Cột ngang xếp hạng", "PivotTable: Rows dim_product[category], Values [Net Sales danh mục dẫn đầu] và [Net Sales danh mục khác]; "
     "sort Largest to Smallest → PivotChart Stacked Bar; Vertical axis → Categories in reverse order", "Chuỗi dẫn đầu xanh, chuỗi khác xám"],
    ["Cột theo tháng có cờ đỏ", "PivotTable: Rows dim_date[month_label] (năm 2020), Values [Δ tháng có vượt ngưỡng] và [Δ tháng khác] "
     "→ PivotChart Stacked Column", "Chuỗi vượt ngưỡng đỏ #D03B3B, chuỗi khác xám"],
    ["Bảng chi tiết", "PivotTable đặt thẳng trên sheet dashboard: Design → Report Layout → Tabular; tắt Grand Total khi không cần; "
     "Value Filters → Top 10; Conditional Formatting (data bar, chữ đỏ)", "Đặt Number Format trong Value Field Settings"],
], widths=[3.6, 9.4, 4.5])
LIST(["Mọi PivotChart: PivotChart Analyze → Field Buttons → **Hide All**; xoá legend khi chỉ một chuỗi; Gap Width 60 %.",
      "PivotTable → Options → bỏ tick **Autofit column widths on update** để bố cục không nhảy khi đổi slicer."])
VIS = {
    "5.5 Sheet Tong_quan": [
        ["Điểm chính", "Ô CUBEVALUE (gộp ô)", "[Điểm chính Tổng quan] & [Điểm chính Động lực]", "B8:Q9"],
        ["5 thẻ KPI", "Ô CUBEVALUE", "Net Revenue · Orders · AOV · Return Rate · Δ Finance (+ Số cửa hàng vượt ngưỡng); dòng so sánh % LM / % SPLY", "B11:Q14"],
        ["Xu hướng 24 tháng", "Line chart thường + bảng phụ", "Net Revenue theo month_label 2019–2020, tháng chọn nhấn", "B16:J32"],
        ["Cầu doanh thu", "Waterfall từ ô CUBEVALUE", "Gross Sales → −Discounts → −Returns → Net Revenue", "K16:Q32"],
        ["Động lực", "2 Line chart thường + bảng phụ", "Orders (nhấn) và AOV (xám) — 2 biểu đồ riêng, cùng trục tháng, không dùng 2 trục y", "B34:J48"],
        ["Đối soát theo tháng", "PivotChart Stacked Column", "[Δ tháng có vượt ngưỡng] + [Δ tháng khác] năm 2020", "K34:Q48"]],
    "5.6 Sheet Doanh_thu": [
        ["Điểm chính", "Ô CUBEVALUE", "[Tiêu đề danh mục] & tỷ trọng online", "B8:Q9"],
        ["4 thẻ KPI", "Ô CUBEVALUE", "Net Sales (+ % SPLY) · Units · Tỷ trọng online · Giảm giá / Gross", "B11:Q14"],
        ["Danh mục", "PivotChart Stacked Bar", "[Net Sales danh mục dẫn đầu] + [Net Sales danh mục khác] theo category", "B16:G34"],
        ["Quốc gia", "PivotChart Bar", "Net Sales theo dim_store[country] (Online là một dòng), sort giảm dần; tô tay thanh đầu", "H16:L34"],
        ["Kênh", "PivotTable", "channel_name; Net Sales, Orders, AOV, Net Sales % SPLY", "M16:Q34"],
        ["Top 10 cửa hàng", "PivotTable Top 10", "store_name; Net Sales (data bar), Orders, AOV; Value Filters → Top 10 theo Net Sales", "B36:Q48"]],
    "5.7 Sheet Doi_soat": [
        ["Điểm chính", "Ô CUBEVALUE", "[Điểm chính Đối soát]", "B8:Q9"],
        ["5 thẻ KPI", "Ô CUBEVALUE", "Net Revenue · Finance Net Revenue · Δ Finance · Δ sau quy đổi · Số cửa hàng vượt ngưỡng (đỏ khi > 0)", "B11:Q14"],
        ["Phân rã Δ", "Waterfall từ ô CUBEVALUE", "Lệch thời điểm online → Δ sau quy đổi (phần cần giải thích) → Δ Finance (total)", "B16:G32"],
        ["Theo tháng", "PivotChart Stacked Column", "[Δ tháng có vượt ngưỡng] + [Δ tháng khác]; không nối slicer Tháng", "H16:Q32"],
        ["Chi tiết cửa hàng", "PivotTable Top 10", "store_code; Net Revenue, Finance Net Revenue, Δ Finance, Δ Finance %, Δ sau quy đổi, Vượt ngưỡng, Δ tuyệt đối; Top 10 theo [Δ tuyệt đối], sort giảm dần", "B34:L50"],
        ["Online", "PivotChart Line", "Filter channel_code = ONLINE; [Δ Finance] (xám) và [Δ sau quy đổi] (nhấn) theo tháng 2020", "M34:Q50"]],
    "5.8 Sheet Hoan_tra": [
        ["Điểm chính", "Ô CUBEVALUE", "[Tiêu đề hoàn trả]; dòng trùng và đơn PENDING đã loại", "B8:Q9"],
        ["4 thẻ KPI", "Ô CUBEVALUE + 1 ô ghi chú", "Returns · Return Rate (điểm % so với SPLY) · Cancelled Orders · Pending Orders; ô “120 dòng trùng đã loại ở Power Query”", "B11:Q14"],
        ["Xu hướng", "Line chart thường + bảng phụ", "Return Rate 24 tháng, tháng chọn nhấn", "B16:J32"],
        ["Lý do", "PivotChart Bar", "Returns theo fact_return[return_reason], sort giảm dần", "K16:Q32"],
        ["Danh mục", "PivotChart Bar", "Return Rate theo category", "B34:J48"],
        ["Huỷ và chờ giao", "PivotChart Line", "Cancelled Orders (nhấn) và Pending Orders (xám) theo tháng; không nối slicer Tháng", "K34:Q48"]],
}
for title, rows_ in VIS.items():
    H3(title)
    TABLE(["Visual", "Đối tượng Excel", "Trường / measure", "Vùng ô"], rows_, widths=[2.8, 3.6, 9.2, 1.9])
P("Cột Vùng ô là gợi ý theo lưới B–Q; giữ cùng vùng cho Điểm chính và KPI ở cả 4 sheet để người xem chuyển sheet "
  "không bị nhảy mắt. Dòng 2–3 của mỗi sheet là tiêu đề trang và dòng phụ (as-of, tháng chọn).")
H3("5.9 Tương tác và hoàn thiện")
LIST(["**Biểu đồ xu hướng không lọc theo tháng.** Pivot nguồn của biểu đồ 24 tháng và biểu đồ theo tháng: Report "
      "Connections bỏ tick slicer Năm và Tháng; lọc khoảng thời gian bằng Filter của chính pivot đó.",
      "**Xem chi tiết.** Bấm đúp một ô giá trị trong PivotTable (Show Details) để Excel đổ ra sheet mới tối đa 1.000 dòng chi "
      "tiết của ô đó — cách dò nguyên nhân một ô cửa hàng × tháng.",
      "**Làm mới.** Data → Refresh All khi đổi SourceFolder hoặc dữ liệu; CUBEVALUE tự cập nhật theo.",
      "**Tiêu đề và đơn vị.** Ô tiêu đề nêu kết luận; dòng phụ ghi đơn vị (USD, nghìn USD, %, điểm %) và phạm vi.",
      "**Chân trang.** Mỗi sheet một dòng chữ xám nhỏ: nguồn, ngày dữ liệu, quy ước Δ và ngưỡng.",
      "**Khoá bố cục (tuỳ chọn).** Review → Protect Sheet, tick Use PivotTable & PivotChart và Edit objects để người xem vẫn "
      "dùng được slicer."])


# ---------------------------------------------------------------- Bước 6
H2("Bước 6. Đối soát số (DÀNH CHO MENTOR — ẩn khi giao đề)")
P(f"Các số dưới đây tính độc lập bằng pandas từ chính file giao học viên và đã khớp với dashboard tham chiếu. "
  f"Tháng báo cáo {mstr(R['report_month'])}, đủ bộ lọc “Tất cả”.", first=True)
H3(f"6.1 KPI tháng {mstr(R['report_month'])}")
def chg(a, b): return pct((a - b) / abs(b)) if b else "–"
krow = lambda name, k, d=0: [name, nf(cur[k], d), nf(lm[k], d), chg(cur[k], lm[k]), nf(sp[k], d), chg(cur[k], sp[k])]
TABLE(["Chỉ tiêu", "12/2020", "11/2020 (LM)", "Δ% LM", "12/2019 (SPLY)", "Δ% SPLY"], [
    krow("Gross Sales (USD)", "gross_sales", 2), krow("Discounts (USD)", "discounts", 2), krow("Net Sales (USD)", "net_sales", 2),
    krow("Returns (USD)", "returns", 2), krow("Net Revenue (USD)", "net_revenue", 2), krow("Orders", "orders"),
    krow("Units", "units"), krow("AOV (USD/đơn)", "aov", 2), krow("ASP (USD/sp)", "asp", 2),
    ["Return Rate", pct(cur["return_rate"], 2), pct(lm["return_rate"], 2), nf((cur["return_rate"] - lm["return_rate"]) * 100, 2) + " điểm %",
     pct(sp["return_rate"], 2), nf((cur["return_rate"] - sp["return_rate"]) * 100, 2) + " điểm %"],
], widths=[3.6, 2.9, 2.9, 1.9, 2.9, 2.0], align="lrrrrr")
fin12 = R["finance_month"][str(R["report_month"])]
P(f"Finance Net Revenue tháng 12/2020 (gồm cả HQ-ADJ {nf(hq20, 2)} USD): {nf(fin12, 2)} USD; chưa gồm HQ-ADJ: "
  f"{nf(fin12 - hq20, 2)} USD. Δ Finance (không gồm HQ-ADJ) = {nf(fin12 - hq20 - cur['net_revenue'], 2)} USD.")
H3("6.2 Lũy kế năm 2020 so với 2019")
yrow = lambda name, k, d=0: [name, nf(ytd[k], d), nf(ytd_sp[k], d), chg(ytd[k], ytd_sp[k])]
TABLE(["Chỉ tiêu", "2020", "2019", "Δ%"], [
    yrow("Gross Sales (USD)", "gross_sales", 2), yrow("Net Sales (USD)", "net_sales", 2), yrow("Returns (USD)", "returns", 2),
    yrow("Net Revenue (USD)", "net_revenue", 2), yrow("Orders", "orders"), yrow("AOV (USD/đơn)", "aov", 2),
    ["Finance Net Revenue (USD, gồm HQ-ADJ)", nf(R["finance_year"]["2020"], 2), nf(R["finance_year"]["2019"], 2),
     chg(R["finance_year"]["2020"], R["finance_year"]["2019"])],
], widths=[6.5, 4, 4, 3], align="lrrr")
H3("6.3 Cầu nối năm 2020: từ dữ liệu đơn đến Finance")
FIG("03_quy_trinh_doi_soat.png", "Hình 5: Quy trình đối soát (nguồn: diagrams/03_quy_trinh_doi_soat.drawio)")
bridge_rows = [["Net Revenue theo dữ liệu đơn (order_date, return_date)", nf(br["orders_net_revenue"], 2)]]
for k, name in CAUSE_VI.items():
    bridge_rows.append([("+ " if br[k] >= 0 else "− ") + name, nf(br[k], 2)])
bridge_rows.append(["= Finance Net Revenue các cửa hàng", nf(br["finance_net"], 2)])
bridge_rows.append(["+ Điều chỉnh HQ-ADJ (31/12/2020)", nf(hq20, 2)])
bridge_rows.append(["= Tổng Finance năm 2020", nf(br["finance_net"] + hq20, 2)])
TABLE(["Khoản", "USD"], bridge_rows, widths=[13, 4.5], align="lr")
P(f"Phần “Online lệch thời điểm” cộng cả năm còn {nf(br['adj_online_timing'], 2)} USD vì đơn đặt cuối 2019 giao đầu 2020 "
  f"nhiều hơn đơn đặt cuối 2020 giao đầu 2021. Đơn huỷ vẫn ghi nhận: {n_canc_2020} đơn trong năm 2020. "
  "Còn lại chưa giải thích = 0.")
H3("6.4 Các ô tháng × cửa hàng vượt ngưỡng năm 2020")
P(f"Có {len(over20)} ô vượt ngưỡng; học viên đạt yêu cầu khi giải thích đúng nguyên nhân chính của mọi ô, "
  "kèm số tiền.", first=True)
TABLE(["Tháng", "Cửa hàng", "Quốc gia", "Đơn hàng", "Finance", "Δ", "Nguyên nhân chính"],
      [[mstr(int(r.month_id)), r.store_code, r.country, nf(r.orders_net_revenue, 0), nf(r.finance_net, 0),
        nf(r.diff_finance_minus_orders, 0), r.cause] for r in over20.itertuples()],
      widths=[2.0, 1.8, 2.7, 2.3, 2.3, 2.1, 4.3], align="lllrrrl")
H3("6.5 Chênh lệch cài sẵn phía Finance (toàn bộ kỳ)")
grp = anom.groupby("type").agg(n=("anomaly_id", "size"), usd=("effect_usd", "sum")).reset_index()
TVI = {"A_CANCELLED_BOOKED": "Đơn huỷ vẫn ghi nhận (K01–K25)", "B_MISSING_UPLOAD": "Finance thiếu 3 ngày (K26)",
       "C_VAT_INCLUDED": "Doanh thu gồm VAT (K27)", "D_DUPLICATE_POSTING": "Bút toán ghi trùng (K28)",
       "E_MANUAL_ADJ": "Điều chỉnh HQ-ADJ (K29–K31)"}
detail = {t: "; ".join(anom[anom.type == t].detail.head(1)) for t in TVI}
TABLE(["Loại", "Số mục", "Ảnh hưởng (USD)", "Mô tả / cách phát hiện"],
      [[TVI[r.type], str(r.n), nf(r.usd, 2), detail[r.type]] for r in grp.itertuples()],
      widths=[4.5, 1.5, 2.8, 8.7], align="lrrl")
P("Nguyên nhân hệ thống (không phải lỗi, nhưng tạo Δ theo ngày/tháng): online ghi nhận theo ngày giao, hoàn tiền hạch "
  "toán trễ 0–3 ngày. Bẫy phía dữ liệu đơn: không bỏ 120 dòng trùng thì Net Sales cao hơn "
  f"{nf(R['systemic']['S3_DUPLICATE_LINES'], 2)} USD (toàn kỳ); tính nhầm đơn PENDING thì cao hơn "
  f"{nf(R['systemic']['S4_PENDING'], 2)} USD.")
H3("6.6 Checklist kiểm tra hành vi")
LIST(["Chọn tháng 06/2020, quốc gia Mỹ: 2 cửa hàng vượt ngưỡng (S54 do VAT, S45 do đơn huỷ được ghi nhận).",
      "Chọn kênh online: “Δ sau quy đổi” gần 0 ở mọi tháng, khác hẳn “Δ gốc”.",
      "Chọn tháng 12/2020: S12 (Pháp) lệch +10.141 USD, đúng bằng một dòng Finance ngày 10/12/2020 bị ghi 2 lần.",
      "Chọn tháng 02/2020: S24 (Đức) Finance thấp hơn 22.169 USD; xuống mức ngày thấy 3 ngày Finance trống.",
      "Bỏ bước Remove duplicates: Net Sales toàn kỳ tăng đúng bằng giá trị 120 dòng trùng.",
      "Đổi slicer Kênh sang POS rồi chọn Quốc gia: biểu đồ online trống, các ô CUBEVALUE so sánh kỳ hiện “–”, không có ô #N/A hay #VALUE!."])

# ---------------------------------------------------------------- Bước 7
H2("Bước 7. Từ dashboard đến báo cáo")
H3("7.1 Mạch viết")
P("Mỗi phần báo cáo đi theo mạch **Câu hỏi → Phân tích → Bằng chứng (hình, số) → Kết luận**. Báo cáo theo đúng 8 phần "
  "của guild_report_template.pdf; phần IV (Main Analysis) dành cho đối soát.", first=True)
H3("7.2 Câu hỏi gợi ý theo trang")
TABLE(["Trang", "Câu hỏi"], [
    ["Tổng quan", "Tháng 12/2020 doanh nghiệp thực sự bán được bao nhiêu? Vì sao Net Revenue giảm mạnh so với 12/2019?"],
    ["Doanh thu", "Kênh nào, danh mục nào, quốc gia nào đóng góp chính? AOV online khác cửa hàng thế nào?"],
    ["Đối soát", "Chênh lệch với Finance là bao nhiêu? Bao nhiêu phần là lệch thời điểm, bao nhiêu là lỗi thật?"],
    ["Hoàn trả & chất lượng", "Tỷ lệ hoàn trả có bất thường ở danh mục nào? Dữ liệu đơn có vấn đề chất lượng gì?"],
], widths=[3.5, 14])
H3("7.3 Finding, Insight, Recommendation, Limitation")
LIST([f"**Finding** (số liệu): “Năm 2020 Finance ghi nhận cao hơn dữ liệu đơn {nf(br['finance_net'] - br['orders_net_revenue'], 0)} USD "
      f"({pct((br['finance_net'] - br['orders_net_revenue']) / br['finance_net'], 2)}).”",
      "**Insight** (ý nghĩa): “Phần lớn chênh lệch đến từ kênh online do khác mốc ghi nhận; sau khi quy về ngày giao, "
      "chênh lệch còn lại tập trung ở 4 lỗi vận hành có thể sửa.”",
      "**Recommendation** (hành động): “Đồng bộ trạng thái huỷ đơn sang kế toán trong ngày; kiểm tra trùng bút toán trước khi khoá sổ.”",
      "**Limitation** (giới hạn): “Dữ liệu là giả lập trên nền GER; VAT giả định 10 % thống nhất; chưa xét tỷ giá.”"])
H3("7.4 Gợi ý sáng tạo (20 điểm)")
LIST(["Sheet chi tiết theo ngày cho một cửa hàng: PivotTable posting_date × [Finance Net Revenue], [Net Revenue], [Δ Finance], lọc bằng slicer cửa hàng.",
      "Waterfall cầu nối Orders → Finance theo nguyên nhân, tách thêm từng cửa hàng ⚠ (ô CUBEVALUE cho từng cửa hàng hoặc bảng nguyên nhân dựng bằng Power Query).",
      "Phân tích khuyến mãi: Net Sales và AOV theo promo_code; giảm giá có kéo số đơn lên không.",
      "Quy đổi về tiền tệ thanh toán bằng fx_rate; phân tích theo vùng khách hàng bằng dim_customer."])
H3("7.5 Sản phẩm nộp")
LIST(["File Nexa_P1_<HoTen>.xlsx: các query Power Query, Data Model, measure DAX và 4 sheet dashboard (+ sheet sáng tạo nếu có). "
      "Mentor mở file, sửa tham số SourceFolder theo máy mình rồi Data → Refresh All.",
      "Báo cáo PDF 8 phần, có bảng giải thích chênh lệch vượt ngưỡng năm 2020 (tháng, cửa hàng, Δ, nguyên nhân, số tiền).",
      "Phụ lục: danh mục measure, ảnh chụp Applied Steps của fact_sales_line và fact_finance, ảnh chụp từng sheet dashboard."])

# ---------------------------------------------------------------- Phụ lục
H2("Phụ lục A. Danh mục measure")
TABLE(["Nhóm", "Measure"], [[t.split("(")[0].strip(), " · ".join(l.split(":=")[0].strip() for l in ls
                                                         if ":=" in l and not l.startswith(" "))] for t, ls in DAXS],
      widths=[4.5, 13])
H2("Phụ lục B. Khi cần sinh lại dữ liệu hoặc sửa sơ đồ")
LIST(["Tải Global_Electronics_Retailer.zip từ thư mục Drive của project vào source/ (file .zip không đưa lên git).",
      "Chạy lần lượt trong _mentor/: generate_dataset.py → build_dashboard.py → build_dashboard_xlsx.py → build_excel_kit.py "
      "→ chup_dashboard.js → build_diagrams.py → build_guideline.py → build_slides.py. Seed cố định nên số liệu không đổi.",
      "build_excel_kit.py sinh excel_build/ (code M, measures.dax) và chạy lại từng bước M bằng pandas để kiểm số dòng; "
      "guideline và slide in lại đúng code đó.",
      "Sửa sơ đồ: mở diagrams/*.drawio bằng draw.io (app.diagrams.net hoặc bản desktop), sửa, lưu; rồi chạy "
      "`python build_diagrams.py --chi-xuat` để xuất lại PDF/PNG mà không ghi đè bản đã sửa tay; chạy lại build_guideline.py để cập nhật hình.",
      "Dashboard Excel không cần sửa M hay DAX khi dữ liệu đổi: Data → Refresh All. Code M và DAX dùng lại được nguyên vẹn "
      "nếu sau này chuyển sang Power BI."])
H2("Phụ lục C. Toàn bộ code Power Query (M)")
P("Đúng nội dung các file trong excel_build/m/, theo thứ tự tạo query.", first=True)
for i, q in enumerate(PQ["order"], 1):
    H3(f"C.{i} {q}")
    P(PQ["queries"][q]["desc"] + " — Nạp: " + PQ["load"].get(q, "Only Create Connection + Add this data to the Data Model") + ".",
      first=True)
    CODE(PQ["queries"][q]["code"])

doc.core_properties.title = "Dashboard Guideline – Sales Project 1"
doc.core_properties.author = "EDA Mentor"
doc.save(OUT)
print("saved", OUT.name)
