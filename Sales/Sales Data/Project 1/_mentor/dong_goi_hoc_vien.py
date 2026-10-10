# -*- coding: utf-8 -*-
"""Gom bộ file giao học viên vào ../hoc_vien/ và đóng gói ../Nexa_Sales_P1_hoc_vien.zip (zip không đưa lên git).

Chạy SAU build_guideline.py và build_slides.py (hai script đó tự sinh bản học viên vào hoc_vien/).
- Chép sơ đồ PDF (diagrams/*.pdf) → hoc_vien/so_do/.
- Ghi hoc_vien/DANH_SACH_FILE.md.
- Kiểm KHÔNG lộ đáp án: không có Instructor Key, hướng dẫn không có Bước 6 bản mentor, slide không còn slide ẩn
  và không còn lời thuyết trình, không có chữ "CHỈ MENTOR" / "DÀNH CHO MENTOR" / "Instructor".
"""
import shutil
import sys
import zipfile
from pathlib import Path

import docx
from pptx import Presentation

HERE = Path(__file__).resolve().parent
P1 = HERE.parent
HV = P1 / "hoc_vien"
ZIP = P1 / "Nexa_Sales_P1_hoc_vien.zip"

DANH_SACH = [
    ("README.docx", "Đề bài Project 1"),
    ("Huong_dan_hoc_vien.docx", "Hướng dẫn từng bước dựng dashboard Excel (Power Query → Mô hình Dữ liệu → DAX → 4 sheet)"),
    ("Slide_Project_1_hoc_vien.pptx", "Slide buổi học"),
    ("Nexa_Sales_P1_Dataset.xlsx", "Dữ liệu đơn hàng, hàng bán bị trả lại, các bảng chiều, bảng map mã Kế toán"),
    ("finance_revenue.xlsx", "Doanh thu phòng Tài chính – Kế toán đã ghi sổ (để đối soát)"),
    ("so_do/", "Sơ đồ pipeline, mô hình dữ liệu, quy trình đối soát, cây KPI, pipeline Power Query (PDF)"),
    ("guild_report_template.pdf", "Khung báo cáo 8 phần"),
]
CAM = ["CHỈ MENTOR", "DÀNH CHO MENTOR", "Instructor", "Bước 6. Đối soát số", "Đáp án cầu nối"]


def main():
    (HV / "so_do").mkdir(exist_ok=True)
    for f in sorted((P1 / "diagrams").glob("*.pdf")):
        shutil.copy2(f, HV / "so_do" / f.name)
    (HV / "DANH_SACH_FILE.md").write_text(
        "# Bộ file học viên — Sales Project 1 (Nexa Home)\n\n| File | Nội dung |\n|---|---|\n" +
        "".join(f"| `{a}` | {b} |\n" for a, b in DANH_SACH) +
        "\nBắt đầu từ `Huong_dan_hoc_vien.docx`. Cần Excel cho Windows (Microsoft 365 hoặc 2016 trở lên) có bổ trợ Power Pivot.\n",
        encoding="utf-8")

    loi = []
    for a, _ in DANH_SACH:
        if not (HV / a.rstrip("/")).exists():
            loi.append(f"thiếu {a}")
    if any("Instructor" in p.name for p in HV.rglob("*")):
        loi.append("có file Instructor Key trong hoc_vien/")
    if any(p.suffix in (".pq", ".dax", ".py") for p in HV.rglob("*")):
        loi.append("có file code (.pq/.dax/.py) trong hoc_vien/ — bài dựng bằng tay")
    d = docx.Document(HV / "Huong_dan_hoc_vien.docx")
    txt = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    loi += [f"hướng dẫn học viên có chữ «{w}»" for w in CAM if w in txt]
    prs = Presentation(str(HV / "Slide_Project_1_hoc_vien.pptx"))
    for i, s in enumerate(prs.slides, 1):
        if s._element.get("show") == "0":
            loi.append(f"slide {i} đang ẩn (học viên bỏ ẩn được)")
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            loi.append(f"slide {i} còn lời thuyết trình")
        st = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        loi += [f"slide {i} có chữ «{w}»" for w in CAM if w in st]
    if loi:
        print("KHÔNG ĐÓNG GÓI – bộ học viên có vấn đề:\n  " + "\n  ".join(loi))
        sys.exit(1)

    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(HV.rglob("*")):
            if p.is_file():
                z.write(p, Path("Nexa_Sales_P1_hoc_vien") / p.relative_to(HV))
    n = sum(1 for p in HV.rglob("*") if p.is_file())
    print(f"OK · {n} file trong hoc_vien/ · {ZIP.name} ({ZIP.stat().st_size / 1e6:.1f} MB, không đưa lên git)")


if __name__ == "__main__":
    main()
