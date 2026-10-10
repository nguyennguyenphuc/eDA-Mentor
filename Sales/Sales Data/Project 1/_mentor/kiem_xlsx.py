# -*- coding: utf-8 -*-
"""Kiểm các lỗi làm Excel báo "We found a problem with some content" (đòi sửa chữa file) trên mọi .xlsx của Project 1.

Không có Excel trên máy sinh tài liệu, nên kiểm thẳng XML theo các lỗi đã gặp hoặc hay gặp:
  1. ô <f> chứa văn bản chứ không phải công thức (chuỗi bắt đầu bằng "=" bị ghi thành công thức);
  2. data validation có formula1/formula2 bắt đầu bằng "=" (sai chuẩn OOXML);
  3. tiêu đề cột ở sheet khác tên cột khai báo trong Excel Table;
  4. tên Table trùng trong cùng file, hoặc Table không có dòng dữ liệu;
  5. chuỗi dài hơn 32.767 ký tự;
  6. thẻ <definedNames/> rỗng;
  7. định dạng số có ký tự ngoài ASCII không đặt trong ngoặc kép (vd dấu "−"), Excel không nhận.
Chạy: python kiem_xlsx.py  → in OK/LỖI từng file, thoát mã 1 nếu có lỗi.
"""
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

P1 = Path(__file__).resolve().parent.parent
# công thức hợp lệ mở đầu bằng: tham chiếu, tên hàm, số, dấu ngoặc, dấu -, +, chuỗi, tên sheet trong nháy
FORMULA_START = re.compile(r"^(?:[A-Za-z_$'(\-+\"0-9\[@!]|_xlfn\.)")


def kiem(path):
    loi = []
    z = zipfile.ZipFile(path)
    wbxml = z.read("xl/workbook.xml").decode("utf-8")
    if "<definedNames/>" in wbxml:
        loi.append("workbook.xml có <definedNames/> rỗng")
    for n in z.namelist():
        if not re.match(r"xl/worksheets/sheet\d+\.xml$", n):
            continue
        x = z.read(n).decode("utf-8")
        for ref, f in re.findall(r'<c r="([A-Z]+\d+)"[^>]*>\s*<f[^>]*>(.*?)</f>', x):
            if not f or not FORMULA_START.match(f) or f.startswith(" "):
                loi.append(f"{n} {ref}: công thức hỏng (thực ra là chữ?) «{f[:40]}»")
        for f in re.findall(r"<formula[12]>(.*?)</formula[12]>", x):
            if f.startswith("="):
                loi.append(f"{n}: data validation formula bắt đầu bằng '=' «{f[:40]}»")
    if "xl/styles.xml" in z.namelist():
        for code in re.findall(r'formatCode="([^"]*)"', z.read("xl/styles.xml").decode("utf-8")):
            bo_ngoac = re.sub(r"&quot;.*?&quot;", "", code)
            if any(ord(ch) > 127 for ch in bo_ngoac):
                loi.append(f"styles.xml: định dạng số có ký tự lạ ngoài ngoặc kép «{code[:40]}»")
    for n in z.namelist():
        if n == "xl/sharedStrings.xml":
            for t in re.findall(r"<t[^>]*>(.*?)</t>", z.read(n).decode("utf-8"), re.S):
                if len(t) > 32767:
                    loi.append("sharedStrings: chuỗi > 32.767 ký tự")
    wb = load_workbook(path)
    ten = []
    for ws in wb:
        for t in ws.tables.values():
            ten.append(t.name.lower())
            a, b = t.ref.split(":")
            r0, c0 = ws[a].row, ws[a].column
            if ws[b].row <= r0:
                loi.append(f"{ws.title}/{t.name}: Table không có dòng dữ liệu")
            hdr = [ws.cell(r0, c0 + i).value for i in range(len(t.tableColumns))]
            names = [tc.name for tc in t.tableColumns]
            if [str(h) for h in hdr] != names:
                loi.append(f"{ws.title}/{t.name}: tiêu đề ≠ tên cột Table")
    if len(ten) != len(set(ten)):
        loi.append("tên Table trùng trong file")
    return loi


if __name__ == "__main__":
    files = sorted(p for p in P1.rglob("*.xlsx") if "~$" not in p.name)
    n_loi = 0
    for p in files:
        loi = kiem(p)
        n_loi += len(loi)
        print(("OK   " if not loi else "LỖI ") + str(p.relative_to(P1)))
        for l in loi[:10]:
            print("     -", l)
    sys.exit(1 if n_loi else 0)
